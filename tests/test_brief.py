from types import SimpleNamespace

from docx import Document

from sugar_core import analysis, brief
from sugar_core import institutions as inst
from sugar_core import activity_coding as ac
from sugar_core.llm_providers import LLMProvider, ProviderProfile
from sugar_core.networks import set_network
from sugar_core.reference_registry import upsert_entity
from test_workbench import bench, prepare, project as make_project


class FakeChat(LLMProvider):
    def __init__(self, text):
        super().__init__(ProviderProfile(id="c", type="local", endpoint="http://localhost:1/v1", model="m", name="Fake"))
        self.text = text

    def chat(self, messages, **kw):
        return SimpleNamespace(text=self.text, model="m", latency_ms=1.0)


def finished_run(tmp_path):
    wb, proj = bench(tmp_path), make_project(tmp_path)
    prepare(wb, proj)
    started = wb.start_run(proj, wait=True)
    run_id = (started.get("run") or started)["run_id"]
    return wb, proj, run_id


def test_brief_is_traceable_and_states_its_limits(tmp_path):
    wb, proj, run_id = finished_run(tmp_path)
    items = wb.items(proj, run_id)
    one = next(i for i in items if i.status in {"collected", "processed"})
    detail = inst.promote(proj, {"name": "Harbor Reading Room", "network": "subject-net", "country": "Exampleland", "city": "Harbor", "latitude": 10.0, "longitude": 20.0, "status": "active"},
                          [(one, "")], actor="ana")
    inst.verify(proj, detail["entity_id"], detail["fields"]["status"]["claims"][0]["claim_id"], "human_verified", actor="vera")
    upsert_entity(proj.workspace, {"name": "Reference Space", "network": "ref-net", "latitude": 10.01, "longitude": 20.01, "country": "Exampleland", "city": "Harbor", "status": "active"},
                  evidence_refs=[{"source_url": "https://ref.example/space"}])
    set_network(proj, "subject-net", role="subject")
    set_network(proj, "ref-net", role="reference")
    ac.code_items(proj, items)
    result = brief.build_brief(wb, proj, run_id)
    md = result["markdown"]
    for heading in ("# Research brief", "## Question and scope", "## What was and was not read", "## Institutions", "## Where the networks meet", "## Review", "## How this was produced", "## Sources"):
        assert heading in md, heading
    assert "do not establish influence" in md and "Harbor Reading Room" in md and "Incomplete: mastodon" in md      # the failed source is stated, not hidden
    assert "[1]" in md and result["references"][0].startswith("https://") and md.count("](") == 0
    assert all(f"{n}. {url}" in md for n, url in enumerate(result["references"], 1))
    assert "verified" in md.lower() or "not verified" in md
    files = brief.write_brief(proj, md)
    root = proj.workspace.root
    assert (root / files["markdown"]).read_text(encoding="utf-8") == md
    doc = Document(str(root / files["docx"]))
    assert any(p.text.startswith("Research brief") for p in doc.paragraphs) and doc.tables


def test_model_summary_keeps_only_sentences_that_cite_real_sources(tmp_path):
    wb, proj, run_id = finished_run(tmp_path)
    one = next(i for i in wb.items(proj, run_id) if i.status in {"collected", "processed"})
    inst.promote(proj, {"name": "Harbor Reading Room", "status": "active"}, [(one, "")], actor="ana")
    text = "One institution is recorded and active [1]. This sentence has no citation. Another claim cites a source that does not exist [99]."
    result = brief.build_brief(wb, proj, run_id, provider=FakeChat(text))
    assert result["model_summary"] and "## Summary (model-assisted)" in result["markdown"]
    assert "One institution is recorded and active [1]." in result["markdown"] and "no citation" not in result["markdown"] and "[99]" not in result["markdown"]
    assert any("dropped" in w for w in result["warnings"])
    plain = brief.build_brief(wb, proj, run_id)
    assert not plain["model_summary"] and "Summary (model-assisted)" not in plain["markdown"]


def test_relevance_filter_bulk_mark_and_themes_over_http_logic(tmp_path):
    wb, proj, run_id = finished_run(tmp_path)
    items = wb.items(proj, run_id)
    summary = analysis.score_run(proj, items, proj.load_plan())
    assert summary["scored"] == len(items) and sum(summary["bands"].values()) == len(items)
    everything = wb.results(proj, run_id, group_by="none", status="all")
    assert everything["relevance_bands"]["unscored"] == 0 and all(r["relevance"] for r in everything["items"])
    likely = wb.results(proj, run_id, group_by="none", relevance="likely")
    assert all(r["relevance"]["band"] == "likely" for r in likely["items"])
    assert len(analysis.themes(items)["themes"]) >= 0
