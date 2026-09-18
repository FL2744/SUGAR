from pathlib import Path

import pytest
from openpyxl import load_workbook

from sugar_core.observation_storage import load_observations, save_observations
from sugar_core.observations import EvidenceReference, ResearchObservation
from sugar_core.state_review import (
    apply_observation_review_workbook,
    apply_observation_review_workbook_file,
    apply_review_workbook,
    export_review_workbook,
)
from sugar_core.state_schema import AnalyticClaim, StateAssessment, SupportAssessment


SOURCE = "https://example.org/evidence"


def observation() -> ResearchObservation:
    return ResearchObservation(
        observation_type="program",
        title="Reviewed program",
        summary="Evidence for a reviewed program.",
        country="example_host_country",
        city="Bishkek",
        evidence=[EvidenceReference(url=SOURCE)],
        verification_state="human_verified",
        reviewer="source-reviewer",
    )


def assessment(obs: ResearchObservation) -> StateAssessment:
    return StateAssessment(
        observation_id=obs.observation_id,
        strategic_audiences=["students"],
        program_domains=["higher_education"],
        sponsor_support=SupportAssessment(
            level="probable",
            bases=["official_sponsor_source"],
            evidence_refs=[SOURCE],
            review_state="ai_triaged",
        ),
        claims=[
            AnalyticClaim(
                statement="The source indicates official sponsorship.",
                claim_type="support_relationship",
                epistemic_status="analytic_assessment",
                evidence_refs=[SOURCE],
                review_state="ai_triaged",
            )
        ],
        review_state="ai_triaged",
    )


def _set_cell(ws, row: int, header: str, value):
    headers = {cell.value: cell.column for cell in ws[1]}
    ws.cell(row=row, column=headers[header]).value = value


def test_review_workbook_can_promote_assessment_support_and_claim_with_named_reviewers(tmp_path: Path):
    obs = observation()
    original = assessment(obs)
    path = Path(export_review_workbook([obs], [original], tmp_path / "review.xlsx"))
    workbook = load_workbook(path)
    assessments = workbook["assessments"]
    _set_cell(assessments, 2, "decision", "human_verified")
    _set_cell(assessments, 2, "reviewer", "Alice Analyst")
    _set_cell(assessments, 2, "support_level_decision", "confirmed")
    _set_cell(assessments, 2, "support_review_decision", "human_verified")
    _set_cell(assessments, 2, "support_reviewer", "Alice Analyst")
    claims = workbook["claims"]
    _set_cell(claims, 2, "decision", "human_verified")
    _set_cell(claims, 2, "reviewer", "Alice Analyst")
    workbook.save(path)
    workbook.close()

    reviewed = apply_review_workbook([original], path)[0]
    assert reviewed.review_state == "human_verified"
    assert reviewed.sponsor_support.level == "confirmed"
    assert reviewed.sponsor_support.review_state == "human_verified"
    assert reviewed.claims[0].review_state == "human_verified"
    assert reviewed.brief_eligible


def test_review_workbook_rejects_verification_without_reviewer(tmp_path: Path):
    obs = observation()
    original = assessment(obs)
    path = Path(export_review_workbook([obs], [original], tmp_path / "review.xlsx"))
    workbook = load_workbook(path)
    _set_cell(workbook["assessments"], 2, "decision", "human_verified")
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="requires a reviewer"):
        apply_review_workbook([original], path)


def test_review_workbook_will_not_verify_influence_without_causal_evidence(tmp_path: Path):
    obs = observation()
    original = StateAssessment(
        observation_id=obs.observation_id,
        observability_level="outcome_evidence",
        claims=[
            AnalyticClaim(
                statement="The program may have shifted participant attitudes.",
                claim_type="influence",
                epistemic_status="hypothesis",
                evidence_refs=[SOURCE],
                review_state="needs_followup",
            )
        ],
        review_state="needs_followup",
    )
    path = Path(export_review_workbook([obs], [original], tmp_path / "review.xlsx"))
    workbook = load_workbook(path)
    _set_cell(workbook["claims"], 2, "decision", "human_verified")
    _set_cell(workbook["claims"], 2, "reviewer", "Alice Analyst")
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="causal_influence_evidence"):
        apply_review_workbook([original], path)


def test_review_workbook_rejects_unknown_taxonomy_edit(tmp_path: Path):
    obs = observation()
    original = assessment(obs)
    path = Path(export_review_workbook([obs], [original], tmp_path / "review.xlsx"))
    workbook = load_workbook(path)
    _set_cell(workbook["assessments"], 2, "strategic_audiences", "students; invented_group")
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="Unsupported strategic_audience"):
        apply_review_workbook([original], path)


def test_review_workbook_round_trips_observation_verification_and_ai_provenance(tmp_path: Path):
    obs = ResearchObservation(
        observation_type="program",
        title="AI-triaged program",
        summary="A public source documents a program.",
        evidence=[EvidenceReference(url=SOURCE)],
    )
    obs.set_ai_triage(
        labels=["education"],
        confidence=0.81,
        provider="arc",
        model="gpt-oss-120b",
        workflow="diplomacy-lab-triage-v1",
        reason="Program activity is explicit in the source.",
    )
    original = assessment(obs)
    path = Path(export_review_workbook([obs], [original], tmp_path / "review.xlsx"))
    workbook = load_workbook(path)
    observations = workbook["observations"]
    headers = {cell.value: cell.column for cell in observations[1]}
    assert observations.cell(2, headers["ai_provider"]).value == "arc"
    assert observations.cell(2, headers["ai_model"]).value == "gpt-oss-120b"
    assert observations.cell(2, headers["ai_workflow"]).value == "diplomacy-lab-triage-v1"
    _set_cell(observations, 2, "decision", "human_verified")
    _set_cell(observations, 2, "reviewer", "Alice Analyst")
    _set_cell(observations, 2, "verification_notes", "Checked against the cited source.")
    workbook.save(path)
    workbook.close()

    reviewed = apply_observation_review_workbook([obs], path)[0]
    assert reviewed.verification_state == "human_verified"
    assert reviewed.reviewer == "Alice Analyst"
    assert reviewed.ai_provider == "arc"
    assert reviewed.ai_model == "gpt-oss-120b"
    assert reviewed.ai_workflow == "diplomacy-lab-triage-v1"


def test_observation_review_rejects_human_verification_without_named_reviewer(tmp_path: Path):
    obs = ResearchObservation(
        observation_type="program",
        summary="A public source documents a program.",
        evidence=[EvidenceReference(url=SOURCE)],
    )
    obs.set_ai_triage(labels=["education"], confidence=0.7, model="test-model")
    path = Path(export_review_workbook([obs], [assessment(obs)], tmp_path / "review.xlsx"))
    workbook = load_workbook(path)
    _set_cell(workbook["observations"], 2, "decision", "human_verified")
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="requires a reviewer"):
        apply_observation_review_workbook([obs], path)


def test_observation_review_file_writes_reviewed_dataset(tmp_path: Path):
    obs = ResearchObservation(
        observation_type="program",
        summary="A public source documents a program.",
        evidence=[EvidenceReference(url=SOURCE)],
    )
    obs.set_ai_triage(labels=["education"], confidence=0.7, model="test-model")
    source = tmp_path / "observations.csv"
    save_observations([obs], source)
    workbook_path = Path(
        export_review_workbook([obs], [assessment(obs)], tmp_path / "review.xlsx")
    )
    workbook = load_workbook(workbook_path)
    _set_cell(workbook["observations"], 2, "decision", "human_verified")
    _set_cell(workbook["observations"], 2, "reviewer", "Alice Analyst")
    workbook.save(workbook_path)
    workbook.close()

    outputs = apply_observation_review_workbook_file(
        source,
        workbook_path,
        tmp_path / "reviewed.csv",
    )
    assert len(outputs) == 3
    restored = load_observations(outputs[0])[0]
    assert restored.verification_state == "human_verified"
    assert restored.reviewer == "Alice Analyst"
