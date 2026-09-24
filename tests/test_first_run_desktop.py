"""Exercise the new-user, no-credential research path through Windows UI actions."""

import json
import os
import sys
import wave
from pathlib import Path

import pytest
from openpyxl import load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")

from PySide6.QtCore import QEventLoop, QSettings, QTimer  # noqa: E402
from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "SUGAR-Windows"))
from app import MainWindow  # noqa: E402
from backend import BackendRunner  # noqa: E402
from sugar_core.observation_storage import load_observations  # noqa: E402
from sugar_core.workspace import SugarWorkspace  # noqa: E402
from sugar_core.workspace_runtime import latest_workspace_artifact_path  # noqa: E402


def test_first_run_import_review_and_verified_handoff(tmp_path, monkeypatch):
    for name in (
        "SUGAR_LLM_API_KEY", "SUGAR_X_BEARER_TOKEN", "SUGAR_BLUESKY_IDENTIFIER",
        "SUGAR_BLUESKY_APP_PASSWORD", "SUGAR_MASTODON_TOKEN", "SUGAR_WEIBO_COOKIE",
    ):
        monkeypatch.delenv(name, raising=False)
    QSettings.setDefaultFormat(QSettings.IniFormat)
    QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, str(tmp_path / "settings"))
    bridge = os.environ.get("SUGAR_E2E_BRIDGE")
    if bridge:
        monkeypatch.setattr(BackendRunner, "_bridge_location", staticmethod(lambda: (bridge, [])))
    else:
        monkeypatch.setattr(
            BackendRunner, "_bridge_location",
            staticmethod(lambda: (sys.executable, [str(ROOT / "sugar_bridge.py")])),
        )
    _app = QApplication.instance() or QApplication([])
    window = MainWindow(smoke=True)
    errors = []
    events = []
    window.runner.error.connect(errors.append)
    window.runner.event.connect(events.append)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: None)

    def step(action):
        del events[:]
        del errors[:]
        loop = QEventLoop()
        def wait_for_idle():
            if (not window.runner.is_running and not window._workspace_operation_queue
                    and not window._workspace_operation_scheduled):
                loop.quit()
            else:
                QTimer.singleShot(20, wait_for_idle)
        QTimer.singleShot(0, wait_for_idle)
        action()
        QTimer.singleShot(90000, loop.quit)
        loop.exec()
        assert not window.runner.is_running, window.activity.log.toPlainText()[-4000:]
        assert not errors, errors
        assert any(event.get("event") == "complete" for event in events), events

    project = tmp_path / "First Project"
    page = window.state_page
    page.research_workspace.setText(str(project))
    page.research_project_name.setText("First Project")
    step(page._research_workspace_open)
    page.research_question.setPlainText("How are public education programs reaching students in Exampleland?")
    page.research_geographies.setText("Exampleland")
    page.research_audiences.setText("students")
    step(page._research_requirement)
    step(page._research_compile)
    assert page.research_strategy_table.rowCount() > 0
    page.research_strategy_reviewer.setText("First User")
    step(page._research_strategy_approve)
    step(page._research_plan)
    assert page.research_plan_table.rowCount() > 0

    source = tmp_path / "partner.csv"
    source.write_text(
        "platform,native_id,canonical_url,original_text\n"
        "partner,1,https://example.org/program,Example Center announced a public education program for students.\n",
        encoding="utf-8",
    )
    page.research_import_file.setText(str(source))
    page.research_import_system.setText("partner-export")
    step(page._research_import)
    step(page._research_prepare_review)
    step(page._research_review_export)

    workbook = project / "state" / "state_review.xlsx"
    assert workbook.is_file()
    review = load_workbook(workbook)
    for sheet_name in ("observations", "assessments"):
        sheet = review[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, headers["decision"]).value = "human_verified"
        sheet.cell(2, headers["reviewer"]).value = "First User"
    review.save(workbook)
    step(page._research_review_apply)
    step(page._research_handoff)
    bundle = Path(page.research_handoff_bundle.text())
    assert bundle.is_dir()
    step(page._research_handoff_verify)
    verification = json.loads((bundle / "verification.json").read_text(encoding="utf-8"))
    manifest = json.loads((bundle / "manifest.json").read_text(encoding="utf-8"))
    assert verification["status"] == "pass"
    assert manifest["counts"]["records"] == 1
    assert manifest["counts"]["observations"] == 1
    assert window.settings_page.secrets()["llm_api_key"] == ""

    intelligence_scroll = window.stack.widget(window.pages["Intelligence"])
    intelligence = intelligence_scroll.widget()
    assert intelligence is not None
    intelligence.quality_workspace.setText(str(project))
    step(intelligence._next_evidence)
    step(intelligence._content_lineage)
    step(intelligence._evidence_graph)
    step(intelligence._robustness)
    intelligence.semantic_query.setText("public education program students")
    step(intelligence._semantic_search)
    intelligence.capture_html.setText(str(source.with_suffix(".html")))
    source.with_suffix(".html").write_text(
        "<html><head><title>Public Program</title></head><body><p>Public education activity for students.</p></body></html>",
        encoding="utf-8",
    )
    intelligence.capture_url.setText("https://example.org/public-program")
    step(intelligence._capture_page)

    media_file = tmp_path / "program-audio.wav"
    with wave.open(str(media_file), "wb") as audio:
        audio.setnchannels(1); audio.setsampwidth(2); audio.setframerate(8000); audio.writeframes(b"\0\0" * 8000 * 5)
    transcript = tmp_path / "program-audio.vtt"
    transcript.write_text("WEBVTT\n\n00:00:02.000 --> 00:00:04.000\nThe speaker describes the program.\n", encoding="utf-8")
    intelligence.media_file.setText(str(media_file))
    intelligence.media_transcript.setText(str(transcript))
    step(intelligence._ingest_media)
    project_workspace = SugarWorkspace.open(project)
    manifest_artifact = project_workspace.latest_artifact("media_manifest")
    assert manifest_artifact is not None
    manifest_path = project_workspace.artifact_absolute_path(manifest_artifact)
    observation_path = latest_workspace_artifact_path(project_workspace, "observations")
    assert observation_path is not None
    observation = load_observations(observation_path)[0]
    intelligence.media_manifest.setText(str(manifest_path))
    intelligence.media_observation_id.setText(observation.observation_id)
    intelligence.media_start.setText("00:02")
    intelligence.media_end.setText("00:04")
    intelligence.media_quote.setText("The speaker describes the program.")
    step(intelligence._attach_media)
    attached_observations = latest_workspace_artifact_path(project_workspace, "observations")
    assert attached_observations is not None
    assert any(item.media_artifact_id for item in load_observations(attached_observations)[0].evidence)
    for filename in (
        "next_evidence_recommendation.json", "content_lineage.json", "temporal_evidence_graph.json",
        "finding_robustness.json", "semantic_evidence_search.json",
    ):
        assert (project / "outputs" / "intelligence" / filename).is_file()
    window.close()
