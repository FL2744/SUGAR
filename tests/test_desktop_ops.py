import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import sugar_bridge
from sugar_core.desktop_ops import DESKTOP_ANALYTIC_OPERATIONS, run_desktop_analytic_operation
from sugar_core.observation_storage import load_observations, save_observations
from sugar_core.observations import EvidenceReference, ResearchObservation
from sugar_core.state_schema import StateAssessment
from sugar_core.state_workflow import load_state_assessments, save_state_assessments
from sugar_core.workspace import SugarWorkspace


def _dataset(tmp_path: Path) -> tuple[Path, Path]:
    observation = ResearchObservation(
        observation_type="program",
        title="Technology workshop",
        summary="A university technology workshop targeted students.",
        observed_at="2026-09-01T12:00:00Z",
        country="example_host_country",
        city="Bishkek",
        institution_name="Example University",
        program_name="Innovation Workshop",
        actors=["Example University"],
        audiences=["students"],
        themes=["technology"],
        evidence=[
            EvidenceReference(
                url="https://example.org/workshop",
                source_type="official_host_source",
                published_at="2026-09-01T12:00:00Z",
            )
        ],
    )
    observations_path = tmp_path / "observations.csv"
    save_observations([observation], observations_path)
    assessments_path = tmp_path / "assessments.jsonl"
    save_state_assessments([StateAssessment(observation_id=observation.observation_id)], assessments_path)
    return observations_path, assessments_path


def test_bridge_advertises_typed_desktop_operations():
    info = sugar_bridge.backend_info()
    assert info["bridge_protocol"] == 3
    assert "state-package" in info["operations"]
    assert "state-map" in info["operations"]
    assert "intel-synthesize" in info["operations"]
    assert "workspace-init" in info["operations"]
    assert "workspace-status" in info["operations"]
    assert "workspace-register" in info["operations"]
    assert "ingest" in info["operations"]
    assert info["collectors"]["wechat"]["known_item"] is True
    for operation in (
        "research-requirement",
        "research-compile",
        "research-strategy-review",
        "research-strategy-update",
        "research-plan",
        "research-plan-review",
        "research-plan-update",
        "research-import",
        "research-prepare-review",
        "research-collect",
        "research-triage",
        "research-feedback",
        "research-handoff",
        "research-handoff-verify",
    ):
        assert operation in info["operations"]
    assert set(DESKTOP_ANALYTIC_OPERATIONS).issubset(info["operations"])


def test_bridge_redacts_runtime_secrets_from_structured_logs(monkeypatch, capsys):
    secret = "super-secret-runtime-token-12345"
    monkeypatch.setenv("SUGAR_LLM_API_KEY", secret)
    sugar_bridge.secrets_from_environment()
    sugar_bridge.emit(
        "error",
        message=f"upstream failure echoed {secret}",
        nested={"detail": secret},
    )
    output = capsys.readouterr().out
    assert secret not in output
    assert output.count("[REDACTED]") == 2
    monkeypatch.delenv("SUGAR_LLM_API_KEY", raising=False)
    sugar_bridge.secrets_from_environment()


def test_desktop_operation_rejects_unlisted_command(tmp_path):
    with pytest.raises(ValueError, match="Unsupported desktop analytic operation"):
        run_desktop_analytic_operation("run-arbitrary-command", {"output_directory": str(tmp_path)})


def test_desktop_templates_and_intelligence_packet(tmp_path):
    us_template = tmp_path / "us_presence.csv"
    outputs = run_desktop_analytic_operation(
        "state-template-us-sites",
        {"output_file": str(us_template)},
    )
    assert outputs == [str(us_template.resolve())]
    assert us_template.is_file()
    assert "american_space" in us_template.read_text(encoding="utf-8-sig")

    observations_path, assessments_path = _dataset(tmp_path)
    packet_path = tmp_path / "packet.json"
    outputs = run_desktop_analytic_operation(
        "intel-packet",
        {
            "observations": str(observations_path),
            "assessments": str(assessments_path),
            "output_file": str(packet_path),
            "case_limit": 10,
        },
    )
    assert outputs == [str(packet_path.resolve())]
    payload = json.loads(packet_path.read_text(encoding="utf-8"))
    assert payload["corpus"]["observations"] == 1
    assert payload["representative_cases"][0]["location"]["country"] == "example_host_country"


def test_state_audit_desktop_operation_writes_json(tmp_path):
    observations_path, assessments_path = _dataset(tmp_path)
    output = tmp_path / "audit.json"
    outputs = run_desktop_analytic_operation(
        "state-audit",
        {
            "observations": str(observations_path),
            "assessments": str(assessments_path),
            "output_file": str(output),
        },
    )
    assert outputs == [str(output.resolve())]
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["observations"] == 1
    assert payload["assessments"] == 1


def test_desktop_state_map_can_resolve_registered_workspace_inputs(tmp_path):
    workspace = SugarWorkspace.create(tmp_path / "project", name="Project")
    observation = ResearchObservation(
        observation_type="event",
        title="Verified event",
        summary="Verified public event.",
        country="example_host_country",
        city="Bishkek",
        latitude=42.8746,
        longitude=74.5698,
        location_basis="reported_city",
        location_confidence=0.8,
        evidence=[EvidenceReference(url="https://example.org/event")],
        verification_state="human_verified",
        reviewer="analyst",
    )
    observations_path = workspace.path_for("observations") / "observations.csv"
    save_observations([observation], observations_path)
    assessments_path = workspace.path_for("state") / "assessments.jsonl"
    save_state_assessments(
        [
            StateAssessment(
                observation_id=observation.observation_id,
                review_state="human_verified",
                reviewer="analyst",
            )
        ],
        assessments_path,
    )
    workspace.register_artifact("observations", observations_path)
    workspace.register_artifact("state_assessments", assessments_path)

    outputs = run_desktop_analytic_operation("state-map", {"workspace": str(workspace.root)})

    html_outputs = [Path(value) for value in outputs if str(value).endswith(".html")]
    assert len(html_outputs) == 1
    assert html_outputs[0].parent == workspace.path_for("maps")
    assert html_outputs[0].is_file()
    assert workspace.latest_artifact("map") is not None


def test_workspace_human_review_round_trip_needs_no_manual_input_paths(tmp_path):
    workspace = SugarWorkspace.create(tmp_path / "project", name="Review Project")
    observation = ResearchObservation(
        observation_type="program",
        title="AI triaged program",
        summary="A public source describes a student program.",
        evidence=[EvidenceReference(url="https://example.org/review")],
    )
    observation.set_ai_triage(
        labels=["education"],
        confidence=0.8,
        provider="arc",
        model="gpt-oss-120b",
        workflow="diplomacy-lab-triage-v1",
    )
    observations_path = workspace.path_for("observations") / "observations.csv"
    save_observations([observation], observations_path)
    assessment = StateAssessment(
        observation_id=observation.observation_id,
        ai_provider="arc",
        ai_model="gpt-oss-120b",
        ai_workflow="state-department-triage-v1",
        review_state="ai_triaged",
    )
    assessments_path = workspace.path_for("state") / "assessments.jsonl"
    save_state_assessments([assessment], assessments_path)
    workspace.register_artifact("observations", observations_path)
    workspace.register_artifact("state_assessments", assessments_path)

    exported = run_desktop_analytic_operation(
        "state-review-export",
        {"workspace": str(workspace.root)},
    )
    workbook_path = Path(exported[0])
    workbook = load_workbook(workbook_path)
    for sheet_name in ("observations", "assessments"):
        sheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, headers["decision"]).value = "human_verified"
        sheet.cell(2, headers["reviewer"]).value = "Analyst One"
    workbook.save(workbook_path)
    workbook.close()

    outputs = run_desktop_analytic_operation(
        "state-review-apply",
        {"workspace": str(workspace.root)},
    )
    reviewed_observations_path = next(
        Path(value) for value in outputs if str(value).endswith("observations_reviewed.csv")
    )
    reviewed_assessments_path = next(
        Path(value) for value in outputs if str(value).endswith("state_reviewed.jsonl")
    )
    reviewed_observation = load_observations(reviewed_observations_path)[0]
    reviewed_assessment = load_state_assessments(reviewed_assessments_path)[0]
    assert reviewed_observation.verification_state == "human_verified"
    assert reviewed_observation.reviewer == "Analyst One"
    assert reviewed_assessment.review_state == "human_verified"
    assert reviewed_assessment.reviewer == "Analyst One"
    assert reviewed_assessment.brief_eligible
