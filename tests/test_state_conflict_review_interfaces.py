from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from sugar_core.desktop_ops import run_desktop_analytic_operation
from sugar_core.observation_storage import save_observations
from sugar_core.observations import EvidenceReference, ResearchObservation
from sugar_core.source_conflicts import SourceClaim, SourceConflict, load_source_conflicts, save_source_conflicts
from sugar_core.state_cli import build_parser
from sugar_core.state_cli import main as state_main
from sugar_core.state_schema import StateAssessment
from sugar_core.state_workflow import save_state_assessments

STATE_URL = "https://educationusa.state.gov/node/421"
OPERATOR_URL = "https://kyrgyzstan.americancouncils.org/edusa"


def _inputs(tmp_path: Path) -> tuple[Path, Path, Path, SourceConflict]:
    observation = ResearchObservation(
        observation_type="program",
        title="Education advising context",
        summary="Conflict review interface regression.",
        country="Kyrgyzstan",
        city="Bishkek",
        evidence=[EvidenceReference(url=STATE_URL, source_type="official_usg_source")],
    )
    observations_path = tmp_path / "observations.csv"
    save_observations([observation], observations_path)

    assessments_path = tmp_path / "assessments.jsonl"
    save_state_assessments(
        [
            StateAssessment(
                observation_id=observation.observation_id,
                strategic_audiences=["students"],
                program_domains=["higher_education"],
            )
        ],
        assessments_path,
    )

    state_claim = SourceClaim(
        statement="EducationUSA Kyrgyzstan is fully online beginning April 1, 2026.",
        source_url=STATE_URL,
        publisher="U.S. Department of State EducationUSA",
        authority_type="official_authority",
        freshness="current",
        effective_date="2026-04-01",
    )
    operator_claim = SourceClaim(
        statement="EducationUSA Kyrgyzstan provides in-person advising in Bishkek.",
        source_url=OPERATOR_URL,
        publisher="American Councils Kyrgyzstan",
        authority_type="official_operator",
    )
    conflict = SourceConflict(
        topic="EducationUSA Kyrgyzstan service topology",
        conflict_type="service_topology",
        status="provisional_treatment",
        claims=[state_claim, operator_claim],
        preferred_claim_id=state_claim.claim_id,
        treatment="Use State topology provisionally.",
        preference_rationale="State supplies an explicit effective date.",
    )
    conflicts_path = Path(save_source_conflicts([conflict], tmp_path / "conflicts.json"))
    return observations_path, assessments_path, conflicts_path, conflict


def _adjudicate(workbook_path: Path, conflict: SourceConflict) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook["source_conflicts"]
    columns = {str(cell.value or ""): cell.column for cell in sheet[1]}
    sheet.cell(2, columns["decision_status"], "human_adjudicated")
    sheet.cell(2, columns["decision_preferred_claim_id"], conflict.claims[0].claim_id)
    sheet.cell(2, columns["decision_treatment"], "Use the reviewed State-directory topology.")
    sheet.cell(2, columns["decision_preference_rationale"], "Analyst reviewed both contradictory public sources.")
    sheet.cell(2, columns["decision_reviewer"], "Analyst One")
    workbook.save(workbook_path)


def test_review_cli_parser_accepts_source_conflict_inputs_and_outputs():
    parser = build_parser()
    export_args = parser.parse_args(
        [
            "review-export",
            "observations.csv",
            "assessments.jsonl",
            "--source-conflicts",
            "conflicts.json",
            "--output",
            "review.xlsx",
        ]
    )
    assert export_args.source_conflicts == "conflicts.json"

    apply_args = parser.parse_args(
        [
            "review-apply",
            "assessments.jsonl",
            "review.xlsx",
            "--source-conflicts",
            "conflicts.json",
            "--source-conflicts-output",
            "reviewed-conflicts.json",
            "--output",
            "reviewed.jsonl",
        ]
    )
    assert apply_args.source_conflicts == "conflicts.json"
    assert apply_args.source_conflicts_output == "reviewed-conflicts.json"


def test_cli_review_round_trip_applies_conflict_adjudication(tmp_path):
    observations, assessments, conflicts, conflict = _inputs(tmp_path)
    workbook = tmp_path / "review.xlsx"
    reviewed_assessments = tmp_path / "reviewed.jsonl"
    reviewed_conflicts = tmp_path / "reviewed-conflicts.json"

    assert (
        state_main(
            [
                "review-export",
                str(observations),
                str(assessments),
                "--source-conflicts",
                str(conflicts),
                "--output",
                str(workbook),
            ]
        )
        == 0
    )
    _adjudicate(workbook, conflict)
    assert (
        state_main(
            [
                "review-apply",
                str(assessments),
                str(workbook),
                "--source-conflicts",
                str(conflicts),
                "--source-conflicts-output",
                str(reviewed_conflicts),
                "--output",
                str(reviewed_assessments),
            ]
        )
        == 0
    )

    assert reviewed_assessments.is_file()
    reviewed = load_source_conflicts(reviewed_conflicts)
    assert reviewed[0].status == "human_adjudicated"
    assert reviewed[0].reviewer == "Analyst One"


def test_cli_review_apply_refuses_to_drop_populated_conflict_decision(tmp_path):
    observations, assessments, conflicts, conflict = _inputs(tmp_path)
    workbook = tmp_path / "review.xlsx"
    state_main(
        [
            "review-export",
            str(observations),
            str(assessments),
            "--source-conflicts",
            str(conflicts),
            "--output",
            str(workbook),
        ]
    )
    _adjudicate(workbook, conflict)

    with pytest.raises(ValueError, match="contains source-conflict decisions"):
        state_main(
            [
                "review-apply",
                str(assessments),
                str(workbook),
                "--output",
                str(tmp_path / "should-not-write.jsonl"),
            ]
        )
    assert not (tmp_path / "should-not-write.jsonl").exists()


def test_desktop_review_round_trip_applies_conflict_adjudication(tmp_path):
    observations, assessments, conflicts, conflict = _inputs(tmp_path)
    workbook = tmp_path / "desktop-review.xlsx"

    export_outputs = run_desktop_analytic_operation(
        "state-review-export",
        {
            "observations": str(observations),
            "assessments": str(assessments),
            "source_conflicts": str(conflicts),
            "output_file": str(workbook),
        },
    )
    assert str(workbook.resolve()) in export_outputs
    _adjudicate(workbook, conflict)

    assessment_output = tmp_path / "desktop-reviewed.jsonl"
    conflict_output = tmp_path / "desktop-reviewed-conflicts.json"
    outputs = run_desktop_analytic_operation(
        "state-review-apply",
        {
            "assessments": str(assessments),
            "workbook": str(workbook),
            "source_conflicts": str(conflicts),
            "output_file": str(assessment_output),
            "source_conflicts_output_file": str(conflict_output),
        },
    )
    assert str(assessment_output.resolve()) in outputs
    assert str(conflict_output.resolve()) in outputs
    reviewed = load_source_conflicts(conflict_output)
    assert reviewed[0].status == "human_adjudicated"
    assert reviewed[0].reviewer == "Analyst One"


def test_desktop_review_apply_refuses_to_drop_populated_conflict_decision(tmp_path):
    observations, assessments, conflicts, conflict = _inputs(tmp_path)
    workbook = tmp_path / "desktop-review.xlsx"
    run_desktop_analytic_operation(
        "state-review-export",
        {
            "observations": str(observations),
            "assessments": str(assessments),
            "source_conflicts": str(conflicts),
            "output_file": str(workbook),
        },
    )
    _adjudicate(workbook, conflict)

    with pytest.raises(ValueError, match="contains source-conflict decisions"):
        run_desktop_analytic_operation(
            "state-review-apply",
            {
                "assessments": str(assessments),
                "workbook": str(workbook),
                "output_file": str(tmp_path / "should-not-write.jsonl"),
            },
        )
    assert not (tmp_path / "should-not-write.jsonl").exists()
