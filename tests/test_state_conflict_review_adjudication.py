from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sugar_core.observations import EvidenceReference, ResearchObservation
from sugar_core.source_conflicts import (
    SourceClaim,
    SourceConflict,
    load_source_conflicts,
    save_source_conflicts,
)
from sugar_core.state_conflict_package import save_state_package_with_conflicts
from sugar_core.state_conflict_review import (
    apply_source_conflict_review_workbook_file,
    export_review_workbook_with_conflict_file,
    workbook_has_source_conflict_decisions,
)
from sugar_core.state_schema import StateAssessment


STATE_URL = "https://educationusa.state.gov/node/421"
OPERATOR_URL = "https://kyrgyzstan.americancouncils.org/edusa"


def _conflict() -> SourceConflict:
    state = SourceClaim(
        statement="EducationUSA Kyrgyzstan is fully online beginning April 1, 2026.",
        source_url=STATE_URL,
        publisher="U.S. Department of State EducationUSA",
        authority_type="official_authority",
        freshness="current",
        effective_date="2026-04-01",
    )
    operator = SourceClaim(
        statement="EducationUSA Kyrgyzstan provides in-person advising in Bishkek.",
        source_url=OPERATOR_URL,
        publisher="American Councils Kyrgyzstan",
        authority_type="official_operator",
        freshness="unknown",
    )
    return SourceConflict(
        topic="EducationUSA Kyrgyzstan service topology",
        conflict_type="service_topology",
        status="provisional_treatment",
        claims=[state, operator],
        preferred_claim_id=state.claim_id,
        treatment="Represent the current service as virtual while preserving the contradictory operator claim.",
        preference_rationale="The State directory supplies an explicit effective date; the operator page does not.",
    )


def _dataset() -> tuple[ResearchObservation, StateAssessment]:
    observation = ResearchObservation(
        observation_type="program",
        title="Education advising context",
        summary="Public-source record used for conflict-review regression coverage.",
        country="Kyrgyzstan",
        city="Bishkek",
        evidence=[
            EvidenceReference(
                url=STATE_URL,
                source_type="official_usg_source",
                collected_at="2026-09-14T00:00:00Z",
            )
        ],
    )
    assessment = StateAssessment(
        observation_id=observation.observation_id,
        strategic_audiences=["students"],
        program_domains=["higher_education"],
    )
    return observation, assessment


def _review_workbook(tmp_path: Path) -> tuple[Path, Path, SourceConflict]:
    observation, assessment = _dataset()
    conflict = _conflict()
    conflicts_path = Path(save_source_conflicts([conflict], tmp_path / "conflicts.json"))
    workbook_path = tmp_path / "review.xlsx"
    export_review_workbook_with_conflict_file(
        [observation],
        [assessment],
        workbook_path,
        source_conflicts_file=conflicts_path,
    )
    return workbook_path, conflicts_path, conflict


def _decision_columns(sheet) -> dict[str, int]:
    return {str(cell.value or ""): cell.column for cell in sheet[1]}


def _write_human_decision(
    workbook_path: Path,
    conflict: SourceConflict,
    *,
    preferred_claim_id: str | None = None,
) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook["source_conflicts"]
    columns = _decision_columns(sheet)
    row = 2
    sheet.cell(row=row, column=columns["decision_status"], value="human_adjudicated")
    sheet.cell(
        row=row,
        column=columns["decision_preferred_claim_id"],
        value=preferred_claim_id or conflict.claims[1].claim_id,
    )
    sheet.cell(
        row=row,
        column=columns["decision_treatment"],
        value="Use the operator topology after human review of the contradictory public records.",
    )
    sheet.cell(
        row=row,
        column=columns["decision_preference_rationale"],
        value="Analyst compared the source claims and selected the operator claim for this reviewed snapshot.",
    )
    sheet.cell(row=row, column=columns["decision_reviewer"], value="Analyst One")
    sheet.cell(
        row=row,
        column=columns["decision_review_note"],
        value="Human adjudication regression test.",
    )
    workbook.save(workbook_path)


def test_conflict_review_export_has_explicit_blank_decision_contract(tmp_path):
    workbook_path, _, conflict = _review_workbook(tmp_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["source_conflicts"]
    columns = _decision_columns(sheet)

    assert sheet.cell(row=2, column=columns["conflict_id"]).value == conflict.conflict_id
    assert conflict.claims[0].claim_id in str(sheet.cell(row=2, column=columns["claim_options"]).value)
    assert conflict.claims[1].claim_id in str(sheet.cell(row=2, column=columns["claim_options"]).value)
    for name in (
        "decision_status",
        "decision_preferred_claim_id",
        "decision_treatment",
        "decision_preference_rationale",
        "decision_reviewer",
        "decision_review_note",
    ):
        assert name in columns
        assert sheet.cell(row=2, column=columns[name]).value is None

    assert workbook_has_source_conflict_decisions(workbook_path) is False


def test_human_adjudication_requires_explicit_decision_fields(tmp_path):
    workbook_path, conflicts_path, _ = _review_workbook(tmp_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["source_conflicts"]
    columns = _decision_columns(sheet)
    sheet.cell(row=2, column=columns["decision_status"], value="human_adjudicated")
    sheet.cell(row=2, column=columns["decision_reviewer"], value="Analyst One")
    workbook.save(workbook_path)

    assert workbook_has_source_conflict_decisions(workbook_path) is True
    with pytest.raises(ValueError, match="requires explicit values"):
        apply_source_conflict_review_workbook_file(
            conflicts_path,
            workbook_path,
            tmp_path / "reviewed.json",
        )


def test_human_adjudication_rejects_unknown_preferred_claim(tmp_path):
    workbook_path, conflicts_path, conflict = _review_workbook(tmp_path)
    _write_human_decision(workbook_path, conflict, preferred_claim_id="srcclaim_not_in_conflict")

    with pytest.raises(ValueError, match="must reference one of its claims"):
        apply_source_conflict_review_workbook_file(
            conflicts_path,
            workbook_path,
            tmp_path / "reviewed.json",
        )


def test_blank_conflict_decision_preserves_provisional_status(tmp_path):
    workbook_path, conflicts_path, conflict = _review_workbook(tmp_path)
    output = Path(
        apply_source_conflict_review_workbook_file(
            conflicts_path,
            workbook_path,
            tmp_path / "reviewed.json",
        )
    )
    reviewed = load_source_conflicts(output)
    assert len(reviewed) == 1
    assert reviewed[0].conflict_id == conflict.conflict_id
    assert reviewed[0].status == "provisional_treatment"
    assert reviewed[0].reviewer == ""


def test_completed_human_adjudication_round_trips_and_clears_package_warning(tmp_path):
    workbook_path, conflicts_path, conflict = _review_workbook(tmp_path)
    _write_human_decision(workbook_path, conflict)

    reviewed_path = Path(
        apply_source_conflict_review_workbook_file(
            conflicts_path,
            workbook_path,
            tmp_path / "reviewed.json",
        )
    )
    reviewed = load_source_conflicts(reviewed_path)
    assert len(reviewed) == 1
    adjudicated = reviewed[0]
    assert adjudicated.status == "human_adjudicated"
    assert adjudicated.preferred_claim_id == conflict.claims[1].claim_id
    assert adjudicated.reviewer == "Analyst One"
    assert adjudicated.review_note == "Human adjudication regression test."
    assert adjudicated.requires_human_review is False

    observation, assessment = _dataset()
    package_dir = tmp_path / "package"
    save_state_package_with_conflicts(
        [observation],
        [assessment],
        package_dir,
        name="reviewed_case",
        source_conflicts=reviewed,
    )
    audit = json.loads((package_dir / "reviewed_case.audit.json").read_text(encoding="utf-8"))
    assert audit["status"] == "pass"
    assert audit["errors"] == 0
    assert audit["warnings"] == 0
    assert audit["source_conflicts"]["human_adjudicated"] == 1
    assert audit["source_conflicts"]["requiring_human_review"] == 0
