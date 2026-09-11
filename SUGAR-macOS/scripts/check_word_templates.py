#!/usr/bin/env python3
"""Exercise Word headers and footers using templates from the frozen backend."""
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import docx.api
import docx.parts.hdrftr
from PyInstaller.archive.readers import CArchiveReader
from PyInstaller.utils.hooks import collect_data_files


def check(binary):
    archive = CArchiveReader(str(binary))
    resources = collect_data_files('docx')
    original_api = docx.api.__file__
    original_hdrftr = docx.parts.hdrftr.__file__
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        for source, destination in resources:
            name = (Path(destination) / Path(source).name).as_posix()
            if name not in archive.toc:
                raise RuntimeError(f'Missing packaged Word resource: {name}')
        for name in archive.toc:
            if not name.startswith('docx/'):
                continue
            path = root / name
            if not path.resolve().is_relative_to(root.resolve()):
                raise RuntimeError(f'Invalid resource path: {name}')
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(archive.extract(name))
        try:
            # These modules resolve default templates relative to __file__.
            # Point them at the frozen resources, not the development install.
            docx.api.__file__ = str(root / 'docx/api.py')
            docx.parts.hdrftr.__file__ = str(root / 'docx/parts/hdrftr.py')
            document = docx.api.Document()
            document.add_paragraph('Analysis packaging regression check')
            document.sections[0].footer.paragraphs[0].text = 'Social Search Posts Analysis'
            document.sections[0].header.paragraphs[0].text = 'Analysis'
            output = io.BytesIO()
            document.save(output)
            with zipfile.ZipFile(output) as saved:
                assert b'Social Search Posts Analysis' in saved.read('word/footer1.xml')
                assert b'Analysis' in saved.read('word/header1.xml')
        finally:
            docx.api.__file__ = original_api
            docx.parts.hdrftr.__file__ = original_hdrftr
    print(f'PASS: {len(resources)} packaged Word resources; document, header, and footer creation')


if __name__ == '__main__':
    check(Path(sys.argv[1]))
