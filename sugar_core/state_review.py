from __future__ import annotations

import json
from dataclasses import asdict
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .observations import ResearchObservation
from .state_schema import (
    NARRATIVE_TAGS,
    OBSERVABILITY_LEVELS,
    PROGRAM_DOMAINS,
    REVIEW_STATES,
    STRATEGIC_AUDIENCES,
    SUPPORT_LEVELS,
    AnalyticClaim,
    StateAssessment,
    SupportAssessment,
)
from .state_workflow import load_state_assessments, save_state_assessments
from .utils import safe_cell


def _clean(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return " ".join(str(value).split())


def _list_text(values: Iterable[str]) -> str:
    return "; ".join(str(value) for value in values)


def _parse_list(value: Any) -> list[str]:
    text = _clean(value)
    if not text:
        return []
    return [part.strip() for part in text.replace("|", ";").split(";") if part.strip()]


def _formula_safe_frame(rows: list[dict[str, Any]]) -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    if not frame.empty:
        for column in frame.columns:
            frame[column] = frame[column].map(lambda value: safe_cell(value, formula_safe=True))
    return frame


def export_review_workbook(
    observations: Iterable[ResearchObservation],
    assessments: Iterable[StateAssessment],
    output_file: str | Path,
) -> str:
    observations = list(observations)
    assessments = list(assessments)
    observation_map = {row.observation_id: row for row in observations}
    target = Path(output_file)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")

    assessment_rows: list[dict[str, Any]] = []
    claim_rows: list[dict[str, Any]] = []
    for assessment in assessments:
        obs = observation_map.get(assessment.observation_id)
        assessment_rows.append(
            {
                "assessment_id": assessment.assessment_id,
                "observation_id": assessment.observation_id,
                "observation_type": obs.observation_type if obs else "",
                "title": obs.title if obs else "",
                "summary": obs.summary if obs else "",
                "country": obs.country if obs else "",
                "city": obs.city if obs else "",
                "primary_source_url": obs.primary_source_url if obs else "",
                "current_review_state": assessment.review_state,
                "decision": "",
                "reviewer": "",
                "review_note": assessment.review_note,
                "current_support_level": assessment.prc_support.level,
                "support_level_decision": "",
                "current_support_review_state": assessment.prc_support.review_state,
                "support_review_decision": "",
                "support_reviewer": "",
                "support_rationale": assessment.prc_support.rationale,
                "support_evidence_refs": _list_text(assessment.prc_support.evidence_refs),
                "strategic_audiences": _list_text(assessment.strategic_audiences),
                "program_domains": _list_text(assessment.program_domains),
                "narrative_tags": _list_text(assessment.narrative_tags),
                "sponsor_entities": _list_text(assessment.sponsor_entities),
                "host_entities": _list_text(assessment.host_entities),
                "partner_entities": _list_text(assessment.partner_entities),
                "delivery_modes": _list_text(assessment.delivery_modes),
                "policy_relevance": _list_text(assessment.policy_relevance),
                "observability_level": assessment.observability_level,
                "analytic_priority": assessment.analytic_priority,
                "us_overlap": assessment.us_overlap.note,
                "claims": len(assessment.claims),
            }
        )
        for claim in assessment.claims:
            claim_rows.append(
                {
                    "claim_id": claim.claim_id,
                    "assessment_id": assessment.assessment_id,
                    "observation_id": assessment.observation_id,
                    "claim_type": claim.claim_type,
                    "statement": claim.statement,
                    "epistemic_status": claim.epistemic_status,
                    "confidence": claim.confidence,
                    "evidence_refs": _list_text(claim.evidence_refs),
                    "current_review_state": claim.review_state,
                    "decision": "",
                    "reviewer": "",
                    "review_note": claim.review_note,
                    "primary_source_url": obs.primary_source_url if obs else "",
                }
            )

    instructions = _formula_safe_frame(
        [
            {"rule": "Purpose", "guidance": "This workbook records human analytic decisions. It does not edit raw source evidence."},
            {"rule": "Evidence", "guidance": "Do not verify a claim unless its evidence_refs identify source evidence attached to the observation."},
            {"rule": "PRC support", "guidance": "Confirmed support requires explicit evidence and support_review_decision=human_verified with a named reviewer."},
            {"rule": "Influence", "guidance": "Do not verify an influence claim from views, likes, comments, attendance, repetition, or proximity alone. Causal influence requires outcome/causal evidence and will still be audited."},
            {"rule": "Anti-U.S./coordination", "guidance": "Use these labels only when the content or relationship is explicit and source-supported."},
            {"rule": "Decision", "guidance": "Use human_verified, rejected, needs_followup, ai_triaged, or unreviewed. A reviewer is required for human_verified/rejected."},
            {"rule": "Taxonomy edits", "guidance": "Semicolon-separated audience/domain/narrative values may be corrected; unknown taxonomy values will fail import rather than silently enter the dataset."},
            {"rule": "Spreadsheet safety", "guidance": "Source-derived text is exported formula-safe so untrusted content cannot execute as an Excel formula when the workbook opens."},
        ]
    )

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        _formula_safe_frame(assessment_rows).to_excel(writer, index=False, sheet_name="assessments")
        _formula_safe_frame(claim_rows).to_excel(writer, index=False, sheet_name="claims")
        instructions.to_excel(writer, index=False, sheet_name="instructions")

    workbook = load_workbook(target)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    decision_fill = PatternFill("solid", fgColor="FFF2CC")
    for sheet_name in ("assessments", "claims", "instructions"):
        ws = workbook[sheet_name]
        ws.freeze_panes = "A2"
        if ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_cells in ws.columns:
            header = str(column_cells[0].value or "")
            width = 18
            if header in {"summary", "statement", "review_note", "support_rationale", "support_evidence_refs", "evidence_refs"}:
                width = 55
            elif header in {"primary_source_url"}:
                width = 42
            elif header in {"strategic_audiences", "program_domains", "narrative_tags", "sponsor_entities", "host_entities", "partner_entities", "policy_relevance", "us_overlap"}:
                width = 32
            ws.column_dimensions[column_cells[0].column_letter].width = width

    decision_values = '"' + ",".join(sorted(REVIEW_STATES)) + '"'
    support_values = '"' + ",".join(sorted(SUPPORT_LEVELS)) + '"'
    observability_values = '"' + ",".join(sorted(OBSERVABILITY_LEVELS)) + '"'
    for sheet_name in ("assessments", "claims"):
        ws = workbook[sheet_name]
        headers = {str(cell.value): cell.column for cell in ws[1]}
        if "decision" in headers:
            validation = DataValidation(type="list", formula1=decision_values, allow_blank=True)
            ws.add_data_validation(validation)
            col = ws.cell(row=1, column=headers["decision"]).column_letter
            validation.add(f"{col}2:{col}{max(2, ws.max_row)}")
            for cell in ws[col][1:]:
                cell.fill = decision_fill
    ws = workbook["assessments"]
    headers = {str(cell.value): cell.column for cell in ws[1]}
    for header, values in (
        ("support_level_decision", support_values),
        ("support_review_decision", decision_values),
        ("observability_level", observability_values),
    ):
        validation = DataValidation(type="list", formula1=values, allow_blank=True)
        ws.add_data_validation(validation)
        col = ws.cell(row=1, column=headers[header]).column_letter
        validation.add(f"{col}2:{col}{max(2, ws.max_row)}")
        if header != "observability_level":
            for cell in ws[col][1:]:
                cell.fill = decision_fill
    workbook.save(target)
    return str(target.resolve())


def _validate_taxonomy(values: list[str], allowed: set[str], field_name: str) -> list[str]:
    normalized = [value.casefold() for value in values]
    invalid = sorted(set(normalized) - allowed)
    if invalid:
        raise ValueError(f"Unsupported {field_name} values in review workbook: {', '.join(invalid)}")
    return normalized


def apply_review_workbook(
    assessments: Iterable[StateAssessment],
    workbook_file: str | Path,
) -> list[StateAssessment]:
    assessments = list(assessments)
    path = Path(workbook_file)
    if not path.is_file():
        raise FileNotFoundError(path)
    assessment_frame = pd.read_excel(path, sheet_name="assessments")
    claim_frame = pd.read_excel(path, sheet_name="claims")
    by_id = {row.assessment_id: row for row in assessments}

    for raw in assessment_frame.to_dict(orient="records"):
        assessment_id = _clean(raw.get("assessment_id"))
        if not assessment_id:
            continue
        assessment = by_id.get(assessment_id)
        if assessment is None:
            raise ValueError(f"Review workbook references unknown assessment_id: {assessment_id}")

        audiences = _parse_list(raw.get("strategic_audiences"))
        domains = _parse_list(raw.get("program_domains"))
        narratives = _parse_list(raw.get("narrative_tags"))
        assessment.strategic_audiences = _validate_taxonomy(audiences, STRATEGIC_AUDIENCES, "strategic_audience")
        assessment.program_domains = _validate_taxonomy(domains, PROGRAM_DOMAINS, "program_domain")
        assessment.narrative_tags = _validate_taxonomy(narratives, NARRATIVE_TAGS, "narrative_tag")
        assessment.sponsor_entities = _parse_list(raw.get("sponsor_entities"))
        assessment.host_entities = _parse_list(raw.get("host_entities"))
        assessment.partner_entities = _parse_list(raw.get("partner_entities"))
        assessment.delivery_modes = _parse_list(raw.get("delivery_modes"))
        assessment.policy_relevance = _parse_list(raw.get("policy_relevance"))

        observability = _clean(raw.get("observability_level")).casefold()
        if observability:
            if observability not in OBSERVABILITY_LEVELS:
                raise ValueError(f"Unsupported observability_level in review workbook: {observability}")
            assessment.observability_level = observability

        decision = _clean(raw.get("decision")).casefold()
        reviewer = _clean(raw.get("reviewer"))
        if decision:
            if decision not in REVIEW_STATES:
                raise ValueError(f"Unsupported assessment decision: {decision}")
            if decision in {"human_verified", "rejected"} and not reviewer:
                raise ValueError(f"Assessment {assessment_id} decision {decision} requires a reviewer.")
            assessment.review_state = decision
            if reviewer:
                assessment.reviewer = reviewer
        note = _clean(raw.get("review_note"))
        if note:
            assessment.review_note = note

        support_level = _clean(raw.get("support_level_decision")).casefold()
        support_review = _clean(raw.get("support_review_decision")).casefold()
        support_reviewer = _clean(raw.get("support_reviewer"))
        if support_level and support_level not in SUPPORT_LEVELS:
            raise ValueError(f"Unsupported support level decision: {support_level}")
        if support_review and support_review not in REVIEW_STATES:
            raise ValueError(f"Unsupported support review decision: {support_review}")
        if support_review in {"human_verified", "rejected"} and not support_reviewer:
            raise ValueError(f"Support review for {assessment_id} requires a reviewer.")
        support_payload = asdict(assessment.prc_support)
        if support_level:
            support_payload["level"] = support_level
        if support_review:
            support_payload["review_state"] = support_review
        if support_reviewer:
            support_payload["reviewer"] = support_reviewer
        rationale = _clean(raw.get("support_rationale"))
        if rationale:
            support_payload["rationale"] = rationale
        assessment.prc_support = SupportAssessment(**support_payload)

    claim_lookup: dict[str, tuple[StateAssessment, int]] = {}
    for assessment in assessments:
        for index, claim in enumerate(assessment.claims):
            claim_lookup[claim.claim_id] = (assessment, index)

    for raw in claim_frame.to_dict(orient="records"):
        claim_id = _clean(raw.get("claim_id"))
        if not claim_id:
            continue
        located = claim_lookup.get(claim_id)
        if located is None:
            raise ValueError(f"Review workbook references unknown claim_id: {claim_id}")
        assessment, index = located
        claim = assessment.claims[index]
        decision = _clean(raw.get("decision")).casefold()
        reviewer = _clean(raw.get("reviewer"))
        note = _clean(raw.get("review_note"))
        payload = asdict(claim)
        if decision:
            if decision not in REVIEW_STATES:
                raise ValueError(f"Unsupported claim decision: {decision}")
            if decision in {"human_verified", "rejected"} and not reviewer:
                raise ValueError(f"Claim {claim_id} decision {decision} requires a reviewer.")
            if claim.claim_type == "influence" and decision == "human_verified" and assessment.observability_level != "causal_influence_evidence":
                raise ValueError(
                    f"Influence claim {claim_id} cannot be human-verified unless observability_level is causal_influence_evidence."
                )
            payload["review_state"] = decision
        if reviewer:
            payload["reviewer"] = reviewer
        if note:
            payload["review_note"] = note
        assessment.claims[index] = AnalyticClaim(**payload)

    return [StateAssessment.from_dict(asdict(row)) for row in assessments]


def apply_review_workbook_file(
    assessments_file: str | Path,
    workbook_file: str | Path,
    output_file: str | Path,
) -> str:
    assessments = load_state_assessments(assessments_file)
    reviewed = apply_review_workbook(assessments, workbook_file)
    return save_state_assessments(reviewed, output_file)
