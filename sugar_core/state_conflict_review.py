from __future__ import annotations

from dataclasses import asdict
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .observations import ResearchObservation
from .source_conflicts import (
    SourceConflict,
    load_source_conflicts,
    save_source_conflicts,
)
from .state_conflict_package import export_review_workbook_with_conflicts
from .state_schema import StateAssessment
from .utils import safe_cell

_CONFLICT_DECISION_STATUS = "human_adjudicated"
_CONFLICT_DECISION_COLUMNS = (
    "decision_status",
    "decision_preferred_claim_id",
    "decision_treatment",
    "decision_preference_rationale",
    "decision_reviewer",
    "decision_review_note",
)


def _clean(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return " ".join(str(value).split())


def _decorate_conflict_review_sheet(
    workbook_file: str | Path,
    conflicts: Iterable[SourceConflict],
) -> str:
    conflicts = list(conflicts)
    if not conflicts:
        return str(Path(workbook_file).expanduser().resolve())

    target = Path(workbook_file).expanduser().resolve()
    workbook = load_workbook(target)
    if "source_conflicts" not in workbook.sheetnames:
        raise ValueError("Conflict-aware review workbook is missing the source_conflicts sheet.")

    sheet = workbook["source_conflicts"]
    headers = {str(cell.value or ""): cell.column for cell in sheet[1]}
    conflict_id_column = headers.get("conflict_id")
    if not conflict_id_column:
        raise ValueError("source_conflicts review sheet is missing conflict_id.")

    start_column = sheet.max_column + 1
    claim_options_column = start_column
    sheet.cell(row=1, column=claim_options_column, value="claim_options")
    decision_start = claim_options_column + 1
    for offset, header in enumerate(_CONFLICT_DECISION_COLUMNS):
        sheet.cell(row=1, column=decision_start + offset, value=header)

    by_id = {conflict.conflict_id: conflict for conflict in conflicts}
    for row_number in range(2, sheet.max_row + 1):
        conflict_id = _clean(sheet.cell(row=row_number, column=conflict_id_column).value)
        conflict = by_id.get(conflict_id)
        if conflict is None:
            continue
        options = " | ".join(
            f"{claim.claim_id} — {claim.publisher or claim.source_label or claim.source_url}: {claim.statement}"
            for claim in conflict.claims
        )
        sheet.cell(
            row=row_number,
            column=claim_options_column,
            value=safe_cell(options, formula_safe=True),
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    decision_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(decision_start, decision_start + len(_CONFLICT_DECISION_COLUMNS)):
        for row_number in range(2, max(2, sheet.max_row) + 1):
            sheet.cell(row=row_number, column=column).fill = decision_fill

    status_validation = DataValidation(
        type="list",
        formula1=f'"{_CONFLICT_DECISION_STATUS}"',
        allow_blank=True,
    )
    sheet.add_data_validation(status_validation)
    status_letter = sheet.cell(row=1, column=decision_start).column_letter
    status_validation.add(f"{status_letter}2:{status_letter}{max(2, sheet.max_row)}")

    widths = {
        "claim_options": 75,
        "decision_status": 22,
        "decision_preferred_claim_id": 34,
        "decision_treatment": 55,
        "decision_preference_rationale": 55,
        "decision_reviewer": 24,
        "decision_review_note": 55,
    }
    header_by_column = {cell.column: str(cell.value or "") for cell in sheet[1]}
    for column, header in header_by_column.items():
        if header in widths:
            sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = widths[header]

    instructions = workbook["instructions"]
    instructions.append(
        [
            "Source-conflict adjudication",
            (
                "To human-adjudicate a conflict, set decision_status=human_adjudicated and explicitly provide "
                "decision_preferred_claim_id, decision_treatment, decision_preference_rationale, and "
                "decision_reviewer. The claim ID must be one of the listed claim_options. Blank decision fields "
                "leave the conflict unchanged. Partial decision rows fail closed. Existing provisional values are "
                "not silently inherited as a human decision."
            ),
        ]
    )
    workbook.save(target)
    return str(target)


def export_review_workbook_with_conflict_file(
    observations: Iterable[ResearchObservation],
    assessments: Iterable[StateAssessment],
    output_file: str | Path,
    *,
    source_conflicts_file: str | Path | None = None,
) -> str:
    conflicts = load_source_conflicts(source_conflicts_file) if source_conflicts_file else []
    output = export_review_workbook_with_conflicts(
        observations,
        assessments,
        output_file,
        source_conflicts=conflicts,
    )
    return _decorate_conflict_review_sheet(output, conflicts)


def workbook_has_source_conflict_decisions(workbook_file: str | Path) -> bool:
    path = Path(workbook_file).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "source_conflicts" not in workbook.sheetnames:
            return False
        sheet = workbook["source_conflicts"]
        headers = {str(cell.value or ""): cell.column for cell in sheet[1]}
        decision_columns = [headers.get(header) for header in _CONFLICT_DECISION_COLUMNS]
        decision_columns = [column for column in decision_columns if column]
        if not decision_columns:
            return False
        for row_number in range(2, sheet.max_row + 1):
            if any(_clean(sheet.cell(row=row_number, column=column).value) for column in decision_columns):
                return True
        return False
    finally:
        workbook.close()


def apply_source_conflict_review_workbook(
    source_conflicts: Iterable[SourceConflict],
    workbook_file: str | Path,
) -> list[SourceConflict]:
    """Apply explicit human source-conflict adjudications from an analyst workbook.

    The original SourceConflict objects remain the source of truth for claims and immutable identity.
    Only dedicated decision_* columns are applied. Blank rows leave conflicts unchanged, while any
    partial adjudication fails rather than inheriting provisional machine/analyst treatment fields.
    """
    conflicts = [SourceConflict(**asdict(conflict)) for conflict in source_conflicts]
    path = Path(workbook_file).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)

    try:
        frame = pd.read_excel(path, sheet_name="source_conflicts", dtype={"conflict_id": str})
    except ValueError as exc:
        raise ValueError("Review workbook is missing the source_conflicts sheet.") from exc

    by_id = {conflict.conflict_id: conflict for conflict in conflicts}
    seen: set[str] = set()

    for raw in frame.to_dict(orient="records"):
        conflict_id = _clean(raw.get("conflict_id"))
        if not conflict_id:
            continue
        if conflict_id in seen:
            raise ValueError(f"Review workbook contains duplicate source conflict: {conflict_id}")
        seen.add(conflict_id)
        conflict = by_id.get(conflict_id)
        if conflict is None:
            raise ValueError(f"Review workbook references unknown source conflict: {conflict_id}")

        decision = {key: _clean(raw.get(key)) for key in _CONFLICT_DECISION_COLUMNS}
        if not any(decision.values()):
            continue

        status = decision["decision_status"].casefold()
        if not status:
            raise ValueError(f"Source conflict {conflict_id} has partial decision fields but no decision_status.")
        if status != _CONFLICT_DECISION_STATUS:
            raise ValueError(
                f"Unsupported source-conflict review decision for {conflict_id}: {status}. "
                f"Only {_CONFLICT_DECISION_STATUS} is accepted from a human review workbook."
            )

        preferred_claim_id = decision["decision_preferred_claim_id"]
        treatment = decision["decision_treatment"]
        rationale = decision["decision_preference_rationale"]
        reviewer = decision["decision_reviewer"]
        missing = [
            label
            for label, value in (
                ("decision_preferred_claim_id", preferred_claim_id),
                ("decision_treatment", treatment),
                ("decision_preference_rationale", rationale),
                ("decision_reviewer", reviewer),
            )
            if not value
        ]
        if missing:
            raise ValueError(
                f"Human adjudication for source conflict {conflict_id} requires explicit values for: "
                + ", ".join(missing)
            )

        claim_ids = {claim.claim_id for claim in conflict.claims}
        if preferred_claim_id not in claim_ids:
            raise ValueError(
                f"decision_preferred_claim_id for source conflict {conflict_id} must reference one of its claims."
            )

        payload = asdict(conflict)
        payload.update(
            {
                "status": "human_adjudicated",
                "preferred_claim_id": preferred_claim_id,
                "treatment": treatment,
                "preference_rationale": rationale,
                "reviewer": reviewer,
                "review_note": decision["decision_review_note"],
            }
        )
        by_id[conflict_id] = SourceConflict(**payload)

    return [by_id[conflict.conflict_id] for conflict in conflicts]


def apply_source_conflict_review_workbook_file(
    source_conflicts_file: str | Path,
    workbook_file: str | Path,
    output_file: str | Path,
) -> str:
    conflicts = load_source_conflicts(source_conflicts_file)
    reviewed = apply_source_conflict_review_workbook(conflicts, workbook_file)
    return save_source_conflicts(reviewed, output_file)
