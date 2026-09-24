from __future__ import annotations

import csv
import hashlib
import html
import json
import shutil
import uuid
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from .models import PostRecord
from .triage_io import load_post_records
from .workspace import SugarWorkspace

RESEARCH_WORKSPACE_SCHEMA_VERSION = "1.0"
RESEARCH_STATE_FILENAME = "sugar-research.json"
SHARE_MANIFEST_FILENAME = "sugar-share-manifest.json"

INSTITUTION_STATUSES = {"active", "closed", "renamed", "relocated", "planned", "unknown"}
SUBPROJECT_STATUSES = {"active", "paused", "complete", "archived"}
LISTENING_POST_CADENCES = {"manual", "hourly", "daily", "weekly", "monthly"}
COLLABORATOR_ROLES = {"owner", "editor", "reviewer", "viewer"}


def _utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _clean(value: Any) -> str:
    return " ".join(str(value or "").split())


def _clean_list(values: Iterable[Any] | str | None) -> list[str]:
    if values is None:
        return []
    if isinstance(values, str):
        text = values.strip()
        if not text:
            return []
        if text.startswith("["):
            try:
                parsed = json.loads(text)
                if isinstance(parsed, list):
                    values = parsed
                else:
                    values = [text]
            except json.JSONDecodeError:
                values = text.replace("|", ";").split(";")
        else:
            values = text.replace("|", ";").split(";")
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = _clean(value)
        key = text.casefold()
        if text and key not in seen:
            result.append(text)
            seen.add(key)
    return result


def _slug(value: str, fallback: str = "item") -> str:
    result = "".join(char.lower() if char.isalnum() else "_" for char in _clean(value))
    result = "_".join(part for part in result.split("_") if part)
    return result[:80] or fallback


def _stable_id(prefix: str, *values: str) -> str:
    raw = "|".join(_clean(value).casefold() for value in values).encode("utf-8")
    return f"{prefix}_" + hashlib.sha256(raw).hexdigest()[:20]


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _float(value: Any) -> float | None:
    text = _clean(value)
    if not text:
        return None
    try:
        result = float(text)
    except (TypeError, ValueError):
        return None
    return result


def _pick(row: dict[str, Any], *names: str) -> Any:
    folded = {str(key).strip().casefold(): value for key, value in row.items()}
    for name in names:
        if name.casefold() in folded:
            return folded[name.casefold()]
    return ""


@dataclass
class Subproject:
    name: str
    subproject_id: str = ""
    description: str = ""
    parent_id: str = ""
    status: str = "active"
    tags: list[str] = field(default_factory=list)
    created_at: str = field(default_factory=_utc_now)
    updated_at: str = field(default_factory=_utc_now)

    def __post_init__(self) -> None:
        self.name = _clean(self.name)
        if not self.name:
            raise ValueError("Subprojects require a name.")
        self.subproject_id = _clean(self.subproject_id) or f"subproject_{uuid.uuid4().hex[:16]}"
        self.description = _clean(self.description)
        self.parent_id = _clean(self.parent_id)
        self.status = _clean(self.status).casefold() or "active"
        if self.status not in SUBPROJECT_STATUSES:
            raise ValueError(f"Unsupported subproject status: {self.status}")
        self.tags = _clean_list(self.tags)


@dataclass
class InstitutionRecord:
    canonical_name: str
    entity_id: str = ""
    entity_type: str = "institution"
    aliases: list[str] = field(default_factory=list)
    country: str = ""
    city: str = ""
    address: str = ""
    latitude: float | None = None
    longitude: float | None = None
    status: str = "unknown"
    opened_at: str = ""
    closed_at: str = ""
    parent_entity_id: str = ""
    successor_entity_id: str = ""
    official_urls: list[str] = field(default_factory=list)
    social_handles: list[str] = field(default_factory=list)
    source_urls: list[str] = field(default_factory=list)
    notes: str = ""
    updated_at: str = field(default_factory=_utc_now)

    def __post_init__(self) -> None:
        self.canonical_name = _clean(self.canonical_name)
        if not self.canonical_name:
            raise ValueError("Institutions require canonical_name.")
        self.entity_type = _clean(self.entity_type).casefold() or "institution"
        self.country = _clean(self.country)
        self.city = _clean(self.city)
        self.address = _clean(self.address)
        self.status = _clean(self.status).casefold() or "unknown"
        if self.status not in INSTITUTION_STATUSES:
            raise ValueError(f"Unsupported institution status: {self.status}")
        self.aliases = _clean_list(self.aliases)
        self.official_urls = _clean_list(self.official_urls)
        self.social_handles = _clean_list(self.social_handles)
        self.source_urls = _clean_list(self.source_urls)
        self.opened_at = _clean(self.opened_at)
        self.closed_at = _clean(self.closed_at)
        self.parent_entity_id = _clean(self.parent_entity_id)
        self.successor_entity_id = _clean(self.successor_entity_id)
        self.notes = _clean(self.notes)
        if self.latitude is not None and not -90 <= float(self.latitude) <= 90:
            raise ValueError("Institution latitude must be between -90 and 90.")
        if self.longitude is not None and not -180 <= float(self.longitude) <= 180:
            raise ValueError("Institution longitude must be between -180 and 180.")
        self.entity_id = _clean(self.entity_id) or _stable_id(
            "institution", self.canonical_name, self.country, self.city
        )

    @property
    def marker_symbol(self) -> str:
        return "☠" if self.status == "closed" else "●"

    @property
    def names(self) -> list[str]:
        return _clean_list([self.canonical_name, *self.aliases])


@dataclass
class ReferenceLayer:
    name: str
    source_path: str
    layer_id: str = ""
    subproject_id: str = ""
    kind: str = "reference"
    visible: bool = True
    file_sha256: str = ""
    row_count: int = 0
    mappable_count: int = 0
    columns: list[str] = field(default_factory=list)
    added_at: str = field(default_factory=_utc_now)

    def __post_init__(self) -> None:
        self.name = _clean(self.name)
        if not self.name:
            raise ValueError("Reference layers require a name.")
        self.layer_id = _clean(self.layer_id) or f"layer_{uuid.uuid4().hex[:16]}"
        self.source_path = _clean(self.source_path)
        if not self.source_path:
            raise ValueError("Reference layers require source_path.")
        self.subproject_id = _clean(self.subproject_id)
        self.kind = _clean(self.kind).casefold() or "reference"
        self.columns = _clean_list(self.columns)


@dataclass
class ListeningPost:
    name: str
    listening_post_id: str = ""
    subproject_id: str = ""
    entity_ids: list[str] = field(default_factory=list)
    handles: list[str] = field(default_factory=list)
    query_terms: list[str] = field(default_factory=list)
    sources: list[str] = field(default_factory=lambda: ["bilibili"])
    cadence: str = "manual"
    enabled: bool = True
    created_at: str = field(default_factory=_utc_now)
    updated_at: str = field(default_factory=_utc_now)
    last_run_at: str = ""
    last_result_count: int = 0
    last_new_count: int = 0
    last_outputs: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.name = _clean(self.name)
        if not self.name:
            raise ValueError("Listening posts require a name.")
        self.listening_post_id = _clean(self.listening_post_id) or f"listen_{uuid.uuid4().hex[:16]}"
        self.subproject_id = _clean(self.subproject_id)
        self.entity_ids = _clean_list(self.entity_ids)
        self.handles = _clean_list(self.handles)
        self.query_terms = _clean_list(self.query_terms)
        self.sources = [value.casefold() for value in _clean_list(self.sources)]
        self.cadence = _clean(self.cadence).casefold() or "manual"
        if self.cadence not in LISTENING_POST_CADENCES:
            raise ValueError(f"Unsupported listening-post cadence: {self.cadence}")
        self.last_outputs = _clean_list(self.last_outputs)


@dataclass
class Collaborator:
    name: str
    collaborator_id: str = ""
    email: str = ""
    role: str = "viewer"
    added_at: str = field(default_factory=_utc_now)

    def __post_init__(self) -> None:
        self.name = _clean(self.name)
        if not self.name:
            raise ValueError("Collaborators require a name.")
        self.email = _clean(self.email)
        self.role = _clean(self.role).casefold() or "viewer"
        if self.role not in COLLABORATOR_ROLES:
            raise ValueError(f"Unsupported collaborator role: {self.role}")
        self.collaborator_id = _clean(self.collaborator_id) or _stable_id(
            "collaborator", self.name, self.email
        )


@dataclass
class SavedView:
    name: str
    saved_view_id: str = ""
    subproject_id: str = ""
    filters: dict[str, Any] = field(default_factory=dict)
    layer_ids: list[str] = field(default_factory=list)
    created_at: str = field(default_factory=_utc_now)

    def __post_init__(self) -> None:
        self.name = _clean(self.name)
        if not self.name:
            raise ValueError("Saved views require a name.")
        self.saved_view_id = _clean(self.saved_view_id) or f"view_{uuid.uuid4().hex[:16]}"
        self.subproject_id = _clean(self.subproject_id)
        self.layer_ids = _clean_list(self.layer_ids)
        if not isinstance(self.filters, dict):
            raise ValueError("Saved-view filters must be an object.")


def _read_rows(path: str | Path) -> list[dict[str, Any]]:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    suffix = source.suffix.casefold()
    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as stream:
            return [dict(row) for row in csv.DictReader(stream)]
    if suffix == ".jsonl":
        rows: list[dict[str, Any]] = []
        for number, line in enumerate(source.read_text(encoding="utf-8-sig").splitlines(), 1):
            if not line.strip():
                continue
            payload = json.loads(line)
            if not isinstance(payload, dict):
                raise ValueError(f"JSONL row {number} must be an object.")
            rows.append(payload)
        return rows
    if suffix in {".json", ".geojson"}:
        payload = json.loads(source.read_text(encoding="utf-8-sig"))
        if isinstance(payload, list):
            return [dict(row) for row in payload if isinstance(row, dict)]
        if isinstance(payload, dict) and payload.get("type") == "FeatureCollection":
            rows = []
            for feature in payload.get("features") or []:
                if not isinstance(feature, dict):
                    continue
                row = dict(feature.get("properties") or {})
                geometry = feature.get("geometry") or {}
                coords = geometry.get("coordinates") or []
                if geometry.get("type") == "Point" and len(coords) >= 2:
                    row.setdefault("longitude", coords[0])
                    row.setdefault("latitude", coords[1])
                rows.append(row)
            return rows
        if isinstance(payload, dict):
            return [payload]
        raise ValueError("JSON reference data must be an object, an array of objects, or a GeoJSON FeatureCollection.")
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator, ())]
            rows = []
            for values in iterator:
                rows.append({headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]})
            return rows
        finally:
            workbook.close()
    raise ValueError("Reference data must be CSV, JSONL, JSON/GeoJSON, or XLSX.")


def _coordinates(row: dict[str, Any]) -> tuple[float, float] | None:
    lat = _float(_pick(row, "latitude", "lat", "y"))
    lon = _float(_pick(row, "longitude", "lon", "lng", "long", "x"))
    if lat is None or lon is None or not (-90 <= lat <= 90 and -180 <= lon <= 180):
        return None
    return lat, lon


def inspect_reference_file(path: str | Path) -> dict[str, Any]:
    rows = _read_rows(path)
    columns = sorted({str(key) for row in rows for key in row})
    return {
        "row_count": len(rows),
        "mappable_count": sum(_coordinates(row) is not None for row in rows),
        "columns": columns,
    }


def institution_from_mapping(row: dict[str, Any]) -> InstitutionRecord:
    name = _clean(_pick(row, "canonical_name", "institution", "institution_name", "name", "center", "centre"))
    if not name:
        raise ValueError("Institution row is missing a name/canonical_name field.")
    closed_at = _clean(_pick(row, "closed_at", "closure_date", "closed", "date_closed"))
    raw_status = _clean(_pick(row, "status", "institution_status")).casefold()
    active_text = _clean(_pick(row, "active", "is_active")).casefold()
    if raw_status not in INSTITUTION_STATUSES:
        if closed_at or active_text in {"false", "0", "no", "inactive", "closed"}:
            raw_status = "closed"
        elif active_text in {"true", "1", "yes", "active"}:
            raw_status = "active"
        else:
            raw_status = "unknown"
    coords = _coordinates(row)
    return InstitutionRecord(
        entity_id=_clean(_pick(row, "entity_id", "institution_id", "id")),
        canonical_name=name,
        entity_type=_clean(_pick(row, "entity_type", "type", "institution_type")) or "institution",
        aliases=_clean_list(_pick(row, "aliases", "alias", "alternate_names")),
        country=_clean(_pick(row, "country", "nation")),
        city=_clean(_pick(row, "city", "locality")),
        address=_clean(_pick(row, "address", "street_address")),
        latitude=coords[0] if coords else None,
        longitude=coords[1] if coords else None,
        status=raw_status,
        opened_at=_clean(_pick(row, "opened_at", "opening_date", "date_opened", "founded")),
        closed_at=closed_at,
        parent_entity_id=_clean(_pick(row, "parent_entity_id", "parent_id")),
        successor_entity_id=_clean(_pick(row, "successor_entity_id", "successor_id")),
        official_urls=_clean_list(_pick(row, "official_urls", "official_url", "website", "websites")),
        social_handles=_clean_list(_pick(row, "social_handles", "handles", "accounts")),
        source_urls=_clean_list(_pick(row, "source_urls", "source_url", "evidence_url", "citation_url")),
        notes=_clean(_pick(row, "notes", "note", "comments")),
    )


class ResearchWorkspaceManager:
    """Portable analyst-facing project memory layered on top of SugarWorkspace.

    The base workspace continues to own identity, artifacts, and local indexing. This
    manager owns human-facing project organization: subprojects, entity lifecycle,
    reference layers, saved searches/history, listening posts, collaborators, and
    reusable views. The state file is portable and deliberately stores no credentials.
    """

    def __init__(self, workspace: SugarWorkspace):
        self.workspace = workspace
        self.state_path = workspace.root / RESEARCH_STATE_FILENAME
        if self.state_path.is_file():
            self.state = json.loads(self.state_path.read_text(encoding="utf-8"))
            self._validate_state()
        else:
            self.state = self._empty_state()
            self._save()

    @classmethod
    def open(cls, path: str | Path) -> "ResearchWorkspaceManager":
        return cls(SugarWorkspace.open(path))

    def _empty_state(self) -> dict[str, Any]:
        return {
            "schema_version": RESEARCH_WORKSPACE_SCHEMA_VERSION,
            "project_id": self.workspace.manifest.project_id,
            "created_at": _utc_now(),
            "updated_at": _utc_now(),
            "subprojects": [],
            "entities": [],
            "reference_layers": [],
            "listening_posts": [],
            "collaborators": [],
            "saved_views": [],
            "history": [],
        }

    def _validate_state(self) -> None:
        if not isinstance(self.state, dict):
            raise ValueError("SUGAR research state must be a JSON object.")
        if str(self.state.get("schema_version") or "") != RESEARCH_WORKSPACE_SCHEMA_VERSION:
            raise ValueError(
                f"Unsupported research workspace schema {self.state.get('schema_version')!r}; "
                f"expected {RESEARCH_WORKSPACE_SCHEMA_VERSION!r}."
            )
        if str(self.state.get("project_id") or "") != self.workspace.manifest.project_id:
            raise ValueError("Research state project_id does not match sugar-project.json.")
        for key in (
            "subprojects",
            "entities",
            "reference_layers",
            "listening_posts",
            "collaborators",
            "saved_views",
            "history",
        ):
            if not isinstance(self.state.get(key), list):
                raise ValueError(f"Research workspace field {key!r} must be a list.")

    def _save(self) -> None:
        self.state["updated_at"] = _utc_now()
        temporary = self.state_path.with_suffix(self.state_path.suffix + ".tmp")
        temporary.write_text(
            json.dumps(self.state, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        temporary.replace(self.state_path)

    def _portable_file_value(self, value: str | Path) -> str:
        path = Path(value).expanduser()
        if not path.is_absolute():
            return path.as_posix()
        stored, external = self.workspace._portable_path(path.resolve())
        return stored if not external else str(path.resolve())

    def _resolve_file_value(self, value: str | Path) -> Path:
        path = Path(value).expanduser()
        return path.resolve() if path.is_absolute() else (self.workspace.root / path).resolve()

    def _append_history(self, event: str, **details: Any) -> dict[str, Any]:
        entry = {
            "history_id": f"history_{uuid.uuid4().hex[:16]}",
            "timestamp": _utc_now(),
            "event": _clean(event),
            **details,
        }
        self.state["history"].append(entry)
        # Keep the portable JSON practical while retaining a substantial project history.
        if len(self.state["history"]) > 10000:
            self.state["history"] = self.state["history"][-10000:]
        return entry

    def status(self) -> dict[str, Any]:
        histories = self.state["history"]
        return {
            "schema_version": RESEARCH_WORKSPACE_SCHEMA_VERSION,
            "project_id": self.workspace.manifest.project_id,
            "name": self.workspace.manifest.name,
            "root": str(self.workspace.root),
            "state_file": str(self.state_path),
            "subprojects": len(self.state["subprojects"]),
            "entities": len(self.state["entities"]),
            "closed_entities": sum(item.get("status") == "closed" for item in self.state["entities"]),
            "reference_layers": len(self.state["reference_layers"]),
            "listening_posts": len(self.state["listening_posts"]),
            "enabled_listening_posts": sum(bool(item.get("enabled", True)) for item in self.state["listening_posts"]),
            "collaborators": len(self.state["collaborators"]),
            "saved_views": len(self.state["saved_views"]),
            "history_events": len(histories),
            "recent_history": histories[-25:],
        }

    def add_subproject(
        self,
        name: str,
        *,
        description: str = "",
        parent_id: str = "",
        status: str = "active",
        tags: Iterable[str] = (),
    ) -> Subproject:
        parent_id = _clean(parent_id)
        if parent_id and not any(item.get("subproject_id") == parent_id for item in self.state["subprojects"]):
            raise ValueError(f"Unknown parent subproject: {parent_id}")
        subproject = Subproject(
            name=name,
            description=description,
            parent_id=parent_id,
            status=status,
            tags=list(tags),
        )
        self.state["subprojects"].append(asdict(subproject))
        self._append_history("subproject_added", subproject_id=subproject.subproject_id, name=subproject.name)
        self._save()
        return subproject

    def add_collaborator(self, name: str, *, email: str = "", role: str = "viewer") -> Collaborator:
        collaborator = Collaborator(name=name, email=email, role=role)
        existing = next(
            (item for item in self.state["collaborators"] if item.get("collaborator_id") == collaborator.collaborator_id),
            None,
        )
        if existing:
            existing.update(asdict(collaborator))
        else:
            self.state["collaborators"].append(asdict(collaborator))
        self._append_history("collaborator_saved", collaborator_id=collaborator.collaborator_id, role=collaborator.role)
        self._save()
        return collaborator

    def save_view(
        self,
        name: str,
        *,
        subproject_id: str = "",
        filters: dict[str, Any] | None = None,
        layer_ids: Iterable[str] = (),
    ) -> SavedView:
        view = SavedView(
            name=name,
            subproject_id=subproject_id,
            filters=filters or {},
            layer_ids=list(layer_ids),
        )
        self.state["saved_views"].append(asdict(view))
        self._append_history("saved_view_added", saved_view_id=view.saved_view_id, name=view.name)
        self._save()
        return view

    def import_institutions(
        self,
        path: str | Path,
        *,
        subproject_id: str = "",
        replace_existing: bool = True,
    ) -> dict[str, Any]:
        source = Path(path).expanduser().resolve()
        rows = _read_rows(source)
        accepted = 0
        rejected: list[dict[str, Any]] = []
        entities_by_id = {item.get("entity_id"): item for item in self.state["entities"]}
        for index, row in enumerate(rows, 1):
            try:
                institution = institution_from_mapping(row)
            except Exception as exc:
                rejected.append({"row": index, "error": str(exc)})
                continue
            payload = asdict(institution)
            payload["subproject_id"] = _clean(subproject_id)
            existing = entities_by_id.get(institution.entity_id)
            if existing is not None:
                if replace_existing:
                    existing.update(payload)
                    accepted += 1
                continue
            self.state["entities"].append(payload)
            entities_by_id[institution.entity_id] = payload
            accepted += 1
        self.workspace.register_artifact(
            "institution_source",
            source,
            label=source.name,
            metadata={
                "subproject_id": _clean(subproject_id),
                "accepted": accepted,
                "rejected": len(rejected),
            },
        )
        self._append_history(
            "institutions_imported",
            source=str(source),
            subproject_id=_clean(subproject_id),
            accepted=accepted,
            rejected=len(rejected),
        )
        self._save()
        return {
            "source": str(source),
            "accepted": accepted,
            "rejected": rejected,
            "total_entities": len(self.state["entities"]),
        }

    def add_reference_layer(
        self,
        path: str | Path,
        *,
        name: str = "",
        subproject_id: str = "",
        kind: str = "reference",
        visible: bool = True,
        copy_into_workspace: bool = True,
    ) -> ReferenceLayer:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        inspection = inspect_reference_file(source)
        layer_id = f"layer_{uuid.uuid4().hex[:16]}"
        stored = source
        if copy_into_workspace:
            directory = self.workspace.path_for("references") / "layers"
            directory.mkdir(parents=True, exist_ok=True)
            destination = directory / f"{layer_id}_{_slug(source.stem)}{source.suffix.casefold()}"
            shutil.copy2(source, destination)
            stored = destination.resolve()
        layer = ReferenceLayer(
            layer_id=layer_id,
            name=name or source.stem,
            source_path=self.workspace._portable_path(stored)[0] if not self.workspace._portable_path(stored)[1] else str(stored),
            subproject_id=subproject_id,
            kind=kind,
            visible=visible,
            file_sha256=_file_sha256(stored),
            row_count=inspection["row_count"],
            mappable_count=inspection["mappable_count"],
            columns=inspection["columns"],
        )
        self.state["reference_layers"].append(asdict(layer))
        self.workspace.register_artifact(
            "reference_layer",
            stored,
            label=layer.name,
            metadata={
                "layer_id": layer.layer_id,
                "subproject_id": layer.subproject_id,
                "kind": layer.kind,
                "row_count": layer.row_count,
                "mappable_count": layer.mappable_count,
            },
        )
        self._append_history(
            "reference_layer_added",
            layer_id=layer.layer_id,
            name=layer.name,
            subproject_id=layer.subproject_id,
            row_count=layer.row_count,
            mappable_count=layer.mappable_count,
        )
        self._save()
        return layer

    def add_listening_post(
        self,
        name: str,
        *,
        subproject_id: str = "",
        entity_ids: Iterable[str] = (),
        handles: Iterable[str] = (),
        query_terms: Iterable[str] = (),
        sources: Iterable[str] = ("bilibili",),
        cadence: str = "manual",
        enabled: bool = True,
    ) -> ListeningPost:
        unknown_entities = [
            entity_id for entity_id in _clean_list(entity_ids)
            if not any(item.get("entity_id") == entity_id for item in self.state["entities"])
        ]
        if unknown_entities:
            raise ValueError(f"Unknown listening-post entities: {', '.join(unknown_entities)}")
        post = ListeningPost(
            name=name,
            subproject_id=subproject_id,
            entity_ids=list(entity_ids),
            handles=list(handles),
            query_terms=list(query_terms),
            sources=list(sources),
            cadence=cadence,
            enabled=enabled,
        )
        self.state["listening_posts"].append(asdict(post))
        self._append_history("listening_post_added", listening_post_id=post.listening_post_id, name=post.name)
        self._save()
        return post

    def listening_post_terms(self, listening_post_id: str) -> list[str]:
        raw = next(
            (item for item in self.state["listening_posts"] if item.get("listening_post_id") == listening_post_id),
            None,
        )
        if raw is None:
            raise ValueError(f"Unknown listening post: {listening_post_id}")
        terms = _clean_list([*(raw.get("query_terms") or []), *(raw.get("handles") or [])])
        entity_ids = set(raw.get("entity_ids") or [])
        for entity in self.state["entities"]:
            if entity.get("entity_id") not in entity_ids:
                continue
            terms = _clean_list([
                *terms,
                entity.get("canonical_name", ""),
                *(entity.get("aliases") or []),
                *(entity.get("social_handles") or []),
            ])
        return terms

    def record_search(
        self,
        *,
        operation: str,
        terms: Iterable[str],
        sources: Iterable[str],
        outputs: Iterable[str] = (),
        result_count: int | None = None,
        subproject_id: str = "",
        listening_post_id: str = "",
        filters: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        entry = self._append_history(
            "search",
            operation=_clean(operation),
            subproject_id=_clean(subproject_id),
            listening_post_id=_clean(listening_post_id),
            terms=_clean_list(terms),
            sources=[value.casefold() for value in _clean_list(sources)],
            filters=filters or {},
            result_count=result_count,
            outputs=[self._portable_file_value(value) for value in _clean_list(outputs)],
        )
        self._save()
        return entry

    def search_history(self, text: str = "", *, subproject_id: str = "") -> list[dict[str, Any]]:
        needle = _clean(text).casefold()
        subproject_id = _clean(subproject_id)
        rows = []
        for item in reversed(self.state["history"]):
            if item.get("event") != "search":
                continue
            if subproject_id and item.get("subproject_id") != subproject_id:
                continue
            haystack = json.dumps(item, ensure_ascii=False).casefold()
            if needle and needle not in haystack:
                continue
            rows.append(item)
        return rows

    def run_listening_post(
        self,
        listening_post_id: str,
        *,
        secrets: dict[str, str] | None = None,
        overrides: dict[str, Any] | None = None,
        progress=None,
    ) -> dict[str, Any]:
        raw = next(
            (item for item in self.state["listening_posts"] if item.get("listening_post_id") == listening_post_id),
            None,
        )
        if raw is None:
            raise ValueError(f"Unknown listening post: {listening_post_id}")
        if not bool(raw.get("enabled", True)):
            raise ValueError("Listening post is disabled.")
        terms = self.listening_post_terms(listening_post_id)
        if not terms:
            raise ValueError("Listening post has no query terms, handles, or linked entities.")
        from .service import run_search

        config: dict[str, Any] = {
            "workspace": str(self.workspace.root),
            "terms": terms,
            "sources": list(raw.get("sources") or ["bilibili"]),
            "translate_posts": False,
            "infer_locations": False,
            "continue_on_source_error": True,
            "subproject_id": str(raw.get("subproject_id") or ""),
            "listening_post_id": listening_post_id,
        }
        if overrides:
            config.update(overrides)
            config["workspace"] = str(self.workspace.root)
            config["terms"] = terms
            config["sources"] = list(raw.get("sources") or ["bilibili"])

        previous_keys: set[str] = set()
        previous_outputs = list(raw.get("last_outputs") or [])
        if previous_outputs:
            previous_path = self._resolve_file_value(previous_outputs[0])
            if previous_path.is_file():
                try:
                    previous_keys = {record.record_key for record in load_post_records(previous_path)}
                except Exception:
                    previous_keys = set()

        outputs = run_search(config, secrets or {}, progress=progress)
        # run_search persists its own search-history event. Reload before mutating the
        # listening-post state so this manager cannot overwrite that concurrent save.
        self.state = json.loads(self.state_path.read_text(encoding="utf-8"))
        self._validate_state()
        raw = next(
            item for item in self.state["listening_posts"]
            if item.get("listening_post_id") == listening_post_id
        )
        current_records: list[PostRecord] = []
        if outputs:
            try:
                current_records = load_post_records(outputs[0])
            except Exception:
                current_records = []
        current_keys = {record.record_key for record in current_records}
        new_keys = sorted(current_keys - previous_keys) if previous_keys else sorted(current_keys)
        delta_path = self.workspace.path_for("raw") / f"{_slug(raw.get('name') or listening_post_id)}.{listening_post_id}.delta.json"
        delta_payload = {
            "schema_version": "1.0",
            "generated_at": _utc_now(),
            "listening_post_id": listening_post_id,
            "previous_record_count": len(previous_keys),
            "current_record_count": len(current_keys),
            "new_record_count": len(new_keys),
            "new_record_keys": new_keys,
            "collection_outputs": [self._portable_file_value(value) for value in outputs],
        }
        delta_path.write_text(json.dumps(delta_payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        self.workspace.register_artifact(
            "listening_post_delta",
            delta_path,
            label=f"{raw.get('name', 'Listening post')} delta",
            metadata={"listening_post_id": listening_post_id, "new_record_count": len(new_keys)},
        )
        raw["last_run_at"] = _utc_now()
        raw["updated_at"] = raw["last_run_at"]
        raw["last_result_count"] = len(current_records)
        raw["last_new_count"] = len(new_keys)
        raw["last_outputs"] = [self._portable_file_value(value) for value in outputs]
        self._append_history(
            "listening_post_run",
            listening_post_id=listening_post_id,
            subproject_id=str(raw.get("subproject_id") or ""),
            result_count=len(current_records),
            new_record_count=len(new_keys),
            outputs=[self._portable_file_value(value) for value in outputs],
            delta=self._portable_file_value(delta_path),
        )
        self._save()
        return {
            "listening_post": raw,
            "terms": terms,
            "outputs": outputs,
            "delta": str(delta_path.resolve()),
            "result_count": len(current_records),
            "new_record_count": len(new_keys),
        }

    def _resolve_layer_path(self, layer: dict[str, Any]) -> Path:
        raw = Path(str(layer.get("source_path") or ""))
        if raw.is_absolute():
            return raw.expanduser().resolve()
        return (self.workspace.root / raw).resolve()

    def create_map(
        self,
        output_file: str | Path | None = None,
        *,
        subproject_id: str = "",
        include_closed: bool = True,
    ) -> list[str]:
        try:
            import folium
        except ImportError as exc:
            raise RuntimeError("folium is required to build the project map.") from exc

        entities = [
            item for item in self.state["entities"]
            if (not subproject_id or item.get("subproject_id") == subproject_id)
            and (include_closed or item.get("status") != "closed")
            and item.get("latitude") is not None
            and item.get("longitude") is not None
        ]
        layers = [
            item for item in self.state["reference_layers"]
            if not subproject_id or item.get("subproject_id") in {"", subproject_id}
        ]
        points: list[tuple[float, float]] = []
        for entity in entities:
            points.append((float(entity["latitude"]), float(entity["longitude"])))
        for layer in layers:
            path = self._resolve_layer_path(layer)
            if not path.is_file():
                continue
            for row in _read_rows(path):
                coords = _coordinates(row)
                if coords:
                    points.append(coords)
        center = [
            sum(point[0] for point in points) / len(points),
            sum(point[1] for point in points) / len(points),
        ] if points else [20.0, 0.0]
        map_obj = folium.Map(location=center, zoom_start=3, control_scale=True)

        institution_group = folium.FeatureGroup(name="Institutions", show=True)
        institution_group.add_to(map_obj)
        for entity in entities:
            status = str(entity.get("status") or "unknown")
            symbol = "☠" if status == "closed" else "●"
            symbol_color = "#8b1e1e" if status == "closed" else "#1d4ed8"
            popup_lines = [
                f"<b>{html.escape(str(entity.get('canonical_name') or 'Institution'))}</b>",
                f"Status: {html.escape(status)}",
            ]
            for label, key in (("City", "city"), ("Country", "country"), ("Opened", "opened_at"), ("Closed", "closed_at")):
                value = _clean(entity.get(key))
                if value:
                    popup_lines.append(f"{label}: {html.escape(value)}")
            if entity.get("source_urls"):
                popup_lines.append("Sources: " + ", ".join(
                    f'<a href="{html.escape(url, quote=True)}">{html.escape(url)}</a>'
                    for url in entity.get("source_urls") or []
                ))
            folium.Marker(
                [float(entity["latitude"]), float(entity["longitude"])],
                tooltip=f"{entity.get('canonical_name')} — {status}",
                popup=folium.Popup("<br>".join(popup_lines), max_width=450),
                icon=folium.DivIcon(
                    html=f'<div style="font-size:20px;color:{symbol_color};font-weight:700">{symbol}</div>'
                ),
            ).add_to(institution_group)

        mapped_reference_points = 0
        for layer in layers:
            path = self._resolve_layer_path(layer)
            if not path.is_file():
                continue
            group = folium.FeatureGroup(name=str(layer.get("name") or path.stem), show=bool(layer.get("visible", True)))
            group.add_to(map_obj)
            for index, row in enumerate(_read_rows(path), 1):
                coords = _coordinates(row)
                if not coords:
                    continue
                mapped_reference_points += 1
                label = _clean(_pick(row, "canonical_name", "name", "institution", "title", "site")) or f"Point {index}"
                folium.CircleMarker(
                    coords,
                    radius=5,
                    weight=1,
                    fill=True,
                    fill_opacity=0.75,
                    tooltip=label,
                    popup=folium.Popup(
                        "<br>".join(
                            f"<b>{html.escape(str(key))}:</b> {html.escape(str(value))}"
                            for key, value in row.items()
                            if _clean(value)
                        )[:8000],
                        max_width=500,
                    ),
                ).add_to(group)

        if points:
            map_obj.fit_bounds([[lat, lon] for lat, lon in points], padding=(24, 24))
        legend = f"""
        <div style="position:fixed;bottom:20px;left:20px;z-index:9999;background:white;
                    border:1px solid #888;padding:10px;max-width:360px;font-size:12px">
          <b>SUGAR project map</b><br>
          Institutions: {len(entities)}; imported reference points: {mapped_reference_points}.<br>
          <span style="color:#1d4ed8">●</span> active/other institution &nbsp;
          <span style="color:#8b1e1e">☠</span> closed institution.<br>
          Lifecycle status is descriptive metadata supported by the project dataset; a marker does not imply influence.
        </div>
        """
        map_obj.get_root().html.add_child(folium.Element(legend))
        folium.LayerControl(collapsed=False).add_to(map_obj)
        target = Path(output_file).expanduser().resolve() if output_file else self.workspace.path_for("maps") / "project_reference_map.html"
        target.parent.mkdir(parents=True, exist_ok=True)
        map_obj.save(str(target))
        metadata = target.with_suffix(target.suffix + ".metadata.json")
        metadata.write_text(
            json.dumps(
                {
                    "schema_version": "1.0",
                    "generated_at": _utc_now(),
                    "project_id": self.workspace.manifest.project_id,
                    "subproject_id": _clean(subproject_id),
                    "include_closed": include_closed,
                    "mapped_institutions": len(entities),
                    "mapped_reference_points": mapped_reference_points,
                    "reference_layers": [
                        {
                            "layer_id": layer.get("layer_id"),
                            "name": layer.get("name"),
                            "visible": layer.get("visible", True),
                        }
                        for layer in layers
                    ],
                    "closed_marker": "☠",
                    "semantics": "reference and institution locations; not a measure of influence or causal effect",
                },
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            ) + "\n",
            encoding="utf-8",
        )
        self.workspace.register_artifact("map", target, label="Project reference map", metadata={"operation": "workspace-map"})
        self.workspace.register_artifact("map_metadata", metadata, label="Project map metadata", metadata={"operation": "workspace-map"})
        self._append_history("project_map_created", output=self._portable_file_value(target), subproject_id=_clean(subproject_id))
        self._save()
        return [str(target), str(metadata)]

    def build_conversation_view(
        self,
        records_file: str | Path,
        *,
        output_file: str | Path | None = None,
        subproject_id: str = "",
    ) -> list[str]:
        records = load_post_records(records_file)
        threads = build_conversation_threads(records)
        if output_file:
            html_path = Path(output_file).expanduser().resolve()
        else:
            html_path = self.workspace.path_for("reports") / "conversation_view.html"
        html_path.parent.mkdir(parents=True, exist_ok=True)
        json_path = html_path.with_suffix(".json")
        json_path.write_text(json.dumps({"threads": threads}, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")

        sections: list[str] = []
        for thread in threads:
            turns = []
            for turn in thread["turns"]:
                indent = min(int(turn.get("depth", 0)), 8) * 24
                speaker = html.escape(turn.get("speaker") or "Unknown")
                platform = html.escape(turn.get("platform") or "")
                timestamp = html.escape(turn.get("published_at") or turn.get("collected_at") or "")
                text = html.escape(turn.get("text") or "").replace("\n", "<br>")
                link = html.escape(turn.get("canonical_url") or "", quote=True)
                turns.append(
                    f'<div class="turn" style="margin-left:{indent}px">'
                    f'<div class="meta"><b>{speaker}</b> <span>{platform}</span> <span>{timestamp}</span></div>'
                    f'<div class="text">{text}</div>'
                    + (f'<a href="{link}">source</a>' if link else "")
                    + "</div>"
                )
            sections.append(
                f'<section><h2>{html.escape(thread["thread_id"])}</h2>'
                f'<p>{len(thread["turns"])} messages · {len(thread["speakers"])} speakers</p>'
                + "".join(turns)
                + "</section>"
            )
        document = """<!doctype html><html><head><meta charset="utf-8"><title>SUGAR Conversation View</title>
<style>
body{font-family:system-ui,-apple-system,Segoe UI,sans-serif;margin:24px;max-width:1100px;color:#172033}
section{border:1px solid #dfe5ef;border-radius:10px;padding:16px;margin:0 0 18px}
.turn{border-left:3px solid #94a3b8;padding:8px 12px;margin-top:8px;background:#f8fafc;border-radius:6px}
.meta{font-size:13px}.meta span{color:#64748b;margin-left:8px}.text{margin:5px 0;white-space:normal}
a{color:#235fa8}
</style></head><body><h1>SUGAR Conversation View</h1>
<p>Handles and speakers remain separate; indentation represents reply ancestry when parent records are available.</p>
""" + "".join(sections) + "</body></html>"
        html_path.write_text(document, encoding="utf-8")
        self.workspace.register_artifact(
            "conversation_view",
            html_path,
            label="Conversation view",
            metadata={"subproject_id": _clean(subproject_id), "threads": len(threads)},
        )
        self.workspace.register_artifact(
            "conversation_data",
            json_path,
            label="Conversation data",
            metadata={"subproject_id": _clean(subproject_id), "threads": len(threads)},
        )
        self._append_history(
            "conversation_view_created",
            records_file=str(Path(records_file).expanduser().resolve()),
            threads=len(threads),
            subproject_id=_clean(subproject_id),
            output=self._portable_file_value(html_path),
        )
        self._save()
        return [str(html_path), str(json_path)]

    def export_share_bundle(self, output_file: str | Path | None = None) -> str:
        target = (
            Path(output_file).expanduser().resolve()
            if output_file
            else self.workspace.path_for("exports") / f"{_slug(self.workspace.manifest.name)}.sugarproject.zip"
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        local_files: dict[str, Path] = {}
        for essential in (self.workspace.manifest_path, self.workspace.catalog_path, self.state_path):
            if essential.is_file():
                local_files[essential.relative_to(self.workspace.root).as_posix()] = essential
        omitted_external: list[dict[str, Any]] = []
        for artifact in self.workspace.list_artifacts():
            path = self.workspace.artifact_absolute_path(artifact)
            if artifact.external:
                omitted_external.append({"kind": artifact.kind, "path": artifact.path, "label": artifact.label})
                continue
            if path.is_file() and path != target:
                relative = path.relative_to(self.workspace.root).as_posix()
                if not relative.startswith(".sugar/"):
                    local_files[relative] = path
        files = [
            {"path": relative, "sha256": _file_sha256(path), "size": path.stat().st_size}
            for relative, path in sorted(local_files.items())
        ]
        manifest = {
            "schema_version": "1.0",
            "created_at": _utc_now(),
            "project_id": self.workspace.manifest.project_id,
            "project_name": self.workspace.manifest.name,
            "files": files,
            "omitted_external_artifacts": omitted_external,
            "notes": "Credentials and the rebuildable local .sugar database/cache are not included.",
        }
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(SHARE_MANIFEST_FILENAME, json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n")
            for relative, path in sorted(local_files.items()):
                archive.write(path, relative)
        self.workspace.register_artifact(
            "project_share",
            target,
            label="Portable shared project",
            metadata={"files": len(files), "omitted_external_artifacts": len(omitted_external)},
        )
        self._append_history("project_shared", output=str(target), files=len(files))
        self._save()
        return str(target)


def build_conversation_threads(records: Iterable[PostRecord]) -> list[dict[str, Any]]:
    values = list(records)
    grouped: dict[str, list[PostRecord]] = {}
    for record in values:
        thread_id = record.conversation_id or record.thread_root_key or record.record_key
        grouped.setdefault(thread_id, []).append(record)
    result: list[dict[str, Any]] = []
    for thread_id, items in grouped.items():
        by_key = {item.record_key: item for item in items}

        def depth(record: PostRecord) -> int:
            seen: set[str] = set()
            current = record
            value = 0
            while current.parent_record_key and current.parent_record_key in by_key:
                if current.parent_record_key in seen:
                    break
                seen.add(current.parent_record_key)
                value += 1
                current = by_key[current.parent_record_key]
            return value

        ordered = sorted(items, key=lambda item: (item.published_at or item.collected_at, item.record_key))
        turns = []
        speakers: list[str] = []
        for record in ordered:
            speaker = record.author_handle or record.author_name or "Unknown"
            if record.author_handle and record.author_name:
                speaker = f"{record.author_name} ({record.author_handle})"
            if speaker not in speakers:
                speakers.append(speaker)
            turns.append(
                {
                    "record_key": record.record_key,
                    "parent_record_key": record.parent_record_key,
                    "depth": depth(record),
                    "speaker": speaker,
                    "author_handle": record.author_handle,
                    "author_name": record.author_name,
                    "platform": record.platform,
                    "published_at": record.published_at,
                    "collected_at": record.collected_at,
                    "text": record.translated_text or record.original_text,
                    "canonical_url": record.canonical_url,
                }
            )
        result.append({"thread_id": thread_id, "speakers": speakers, "turns": turns})
    return sorted(result, key=lambda item: (-len(item["turns"]), item["thread_id"]))


def record_workspace_search(
    workspace: SugarWorkspace | None,
    config: dict[str, Any],
    outputs: Iterable[str],
    *,
    operation: str = "search",
    result_count: int | None = None,
) -> None:
    if workspace is None:
        return
    manager = ResearchWorkspaceManager(workspace)
    manager.record_search(
        operation=operation,
        terms=config.get("terms") or [],
        sources=config.get("sources") or [],
        outputs=outputs,
        result_count=result_count,
        subproject_id=str(config.get("subproject_id") or ""),
        listening_post_id=str(config.get("listening_post_id") or ""),
        filters={
            "since": config.get("since") or "",
            "until": config.get("until") or "",
            "max_posts_per_query": config.get("max_posts_per_query"),
            "max_pages_per_query": config.get("max_pages_per_query"),
            "research_requirement_id": config.get("research_requirement_id") or "",
            "plan_branch_ids": config.get("plan_branch_ids") or [],
        },
    )


def import_project_bundle(
    bundle_file: str | Path,
    destination: str | Path,
    *,
    overwrite: bool = False,
) -> SugarWorkspace:
    bundle = Path(bundle_file).expanduser().resolve()
    if not bundle.is_file():
        raise FileNotFoundError(bundle)
    target = Path(destination).expanduser().resolve()
    if target.exists() and any(target.iterdir()) and not overwrite:
        raise FileExistsError(f"Destination is not empty: {target}")
    target.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(bundle) as archive:
        try:
            manifest = json.loads(archive.read(SHARE_MANIFEST_FILENAME))
        except KeyError as exc:
            raise ValueError("Project bundle is missing sugar-share-manifest.json.") from exc
        files = manifest.get("files") or []
        if not isinstance(files, list):
            raise ValueError("Project share manifest files must be a list.")
        for entry in files:
            relative = Path(str(entry.get("path") or ""))
            if not str(relative) or relative.is_absolute() or ".." in relative.parts:
                raise ValueError(f"Unsafe project bundle path: {relative}")
            destination_path = (target / relative).resolve()
            try:
                destination_path.relative_to(target)
            except ValueError as exc:
                raise ValueError(f"Project bundle path escapes destination: {relative}") from exc
            destination_path.parent.mkdir(parents=True, exist_ok=True)
            if destination_path.exists() and not overwrite:
                raise FileExistsError(destination_path)
            with archive.open(relative.as_posix()) as source, destination_path.open("wb") as output:
                shutil.copyfileobj(source, output)
            expected = _clean(entry.get("sha256"))
            if expected and _file_sha256(destination_path) != expected:
                raise ValueError(f"Project bundle hash mismatch: {relative}")
    workspace = SugarWorkspace.open(target)
    ResearchWorkspaceManager(workspace)
    return workspace


def user_support_payload() -> dict[str, Any]:
    return {
        "project_model": [
            "Use one workspace per research effort; create subprojects for country, institution, or thematic tracks.",
            "Searches performed inside a workspace are retained in project history with terms, sources, filters, result counts, and outputs.",
            "Reference layers and institution lifecycle records are portable project data, not map-only decoration.",
        ],
        "listening_posts": [
            "A listening post is a saved monitoring definition built from terms, handles, and linked institutions.",
            "Running it produces ordinary provenance-preserving SUGAR collection artifacts plus a delta of newly observed record identities.",
            "Cadence is stored as project intent; desktop/CLI runs remain explicit and respect each platform's authorized-access boundaries.",
        ],
        "mapping": [
            "Import CSV, JSONL, GeoJSON, or XLSX reference layers. Coordinate-bearing rows become independently toggleable map layers.",
            "Institution records preserve active/closed/renamed/relocated lifecycle state. Closed institutions render with a skull marker on the project map.",
            "Map presence and proximity are descriptive; SUGAR does not convert them into an influence score.",
        ],
        "sharing": [
            "Share Project exports a portable .sugarproject.zip containing registered project-local artifacts and integrity hashes.",
            "Runtime credentials, the rebuildable local SQLite index, cache files, and external artifacts are excluded.",
        ],
        "conversations": [
            "Conversation view groups records by platform conversation/thread identity and keeps each handle/speaker distinct.",
            "Reply ancestry is shown by indentation when parent-record links are present.",
        ],
    }
