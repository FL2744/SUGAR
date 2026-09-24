from __future__ import annotations

import json
import os
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from .workspace import MANIFEST_FILENAME, SugarWorkspace


_SECRET_KEYS = ("api_key", "password", "cookie", "authorization", "secret", "credential", "bearer_token", "app_password")


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _safe_value(value: Any, key: str = "") -> Any:
    if any(token in key.casefold() for token in _SECRET_KEYS):
        return "[omitted]"
    if isinstance(value, dict):
        return {str(child_key): _safe_value(child, str(child_key)) for child_key, child in value.items() if not any(token in str(child_key).casefold() for token in _SECRET_KEYS)}
    if isinstance(value, (list, tuple)):
        return [_safe_value(item) for item in value]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def history_path(workspace: SugarWorkspace) -> Path:
    target = workspace.internal_path / "project-history.jsonl"
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def log_project_event(
    workspace: SugarWorkspace,
    event_type: str,
    payload: dict[str, Any] | None = None,
    *,
    actor: str = "desktop analyst",
) -> dict[str, Any]:
    row = {
        "event_id": str(uuid.uuid4()), "event_type": str(event_type), "project_id": workspace.manifest.project_id,
        "occurred_at": utc_now(), "actor": str(actor or "analyst"), "details": _safe_value(payload or {}),
    }
    with history_path(workspace).open("a", encoding="utf-8", newline="\n") as stream:
        stream.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")
    workspace.register_artifact("project_history", history_path(workspace), label="Project activity and search history")
    return row


def _result_counts(outputs: Iterable[str | Path]) -> list[dict[str, Any]]:
    counts: list[dict[str, Any]] = []
    seen: set[str] = set()
    for raw in outputs:
        path = Path(raw).expanduser()
        if not path.is_file() or path.suffix.casefold() not in {".csv", ".xlsx", ".jsonl", ".ndjson"}:
            continue
        key = str(path.resolve().with_suffix(""))
        if key in seen and path.suffix.casefold() in {".xlsx", ".ndjson"}:
            continue
        seen.add(key)
        try:
            if path.suffix.casefold() == ".csv":
                with path.open("r", encoding="utf-8-sig", newline="") as stream:
                    count = max(0, sum(1 for _ in stream) - 1)
            elif path.suffix.casefold() == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    sheet = workbook.active
                    count = max(0, sheet.max_row - 1)
                finally:
                    workbook.close()
            else:
                with path.open("r", encoding="utf-8-sig") as stream:
                    count = sum(1 for line in stream if line.strip())
            counts.append({"path": path.name, "records": count})
        except (OSError, ValueError, KeyError):
            counts.append({"path": path.name, "records": None, "count_error": True})
    return counts


def _run_evidence(outputs: Iterable[str | Path]) -> dict[str, Any]:
    """Capture the effective query and source coverage emitted by collectors."""
    metadata: dict[str, Any] = {}
    coverage: dict[str, Any] = {}
    for raw in outputs:
        path = Path(raw).expanduser()
        if not path.is_file():
            continue
        suffix = path.suffix.casefold()
        if suffix == ".json" and path.name.endswith(".coverage.json"):
            try:
                value = json.loads(path.read_text(encoding="utf-8-sig"))
                if isinstance(value, dict):
                    coverage = value
            except (OSError, ValueError):
                pass
        elif suffix == ".json" and path.name.endswith(".metadata.json"):
            try:
                value = json.loads(path.read_text(encoding="utf-8-sig"))
                if isinstance(value, dict):
                    metadata = value
            except (OSError, ValueError):
                pass
    return {
        "effective_terms": metadata.get("terms") or coverage.get("terms") or [],
        "effective_sources": metadata.get("sources") or [
            str(key) for key in (coverage.get("sources") or {}).keys()
        ],
        "date_window": {
            "since": metadata.get("since") or coverage.get("since"),
            "until": metadata.get("until") or coverage.get("until"),
        },
        "source_coverage": _safe_value(metadata.get("source_coverage") or coverage),
        "collector_capabilities": _safe_value(metadata.get("collector_capabilities") or {}),
    }


def record_project_run(
    workspace: SugarWorkspace | None,
    *,
    command: str,
    config: dict[str, Any],
    outputs: Iterable[str | Path] = (),
    status: str = "succeeded",
    error: str = "",
    run_id: str | None = None,
    started_at: str | None = None,
) -> dict[str, Any] | None:
    if workspace is None:
        return None
    output_values = list(outputs)
    output_paths = [str(Path(value).expanduser()) for value in output_values]
    evidence = _run_evidence(output_values)
    safe_config = _safe_value(config)
    if isinstance(safe_config, dict):
        safe_config = {**safe_config, **{key: value for key, value in evidence.items() if value not in (None, [], {})}}
    run = {
        "run_id": run_id or str(uuid.uuid4()), "command": str(command),
        "started_at": started_at or utc_now(), "completed_at": utc_now(),
        "status": status, "parameters": safe_config,
        "outputs": output_paths,
        "result_counts": _result_counts(output_values), "error": str(error or ""),
    }
    return log_project_event(workspace, "research_run", run)


def list_project_history(workspace: SugarWorkspace, *, limit: int = 200) -> list[dict[str, Any]]:
    path = history_path(workspace)
    if not path.is_file():
        return []
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig") as stream:
        for line_number, line in enumerate(stream, start=1):
            if not line.strip():
                continue
            item = json.loads(line)
            if not isinstance(item, dict):
                raise ValueError(f"Project history line {line_number} must contain an object.")
            rows.append(item)
    return rows[-max(1, min(10000, limit)):][::-1]


def _relation_path(workspace: SugarWorkspace) -> Path:
    return workspace.internal_path / "project-relations.json"


def _read_relation(workspace: SugarWorkspace) -> dict[str, Any]:
    path = _relation_path(workspace)
    if not path.is_file():
        return {"schema_version": "1.0", "parent": None, "children": []}
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Project relation record must be an object.")
    payload.setdefault("parent", None)
    payload.setdefault("children", [])
    return payload


def _write_relation(workspace: SugarWorkspace, payload: dict[str, Any]) -> None:
    path = _relation_path(workspace)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    temporary.replace(path)
    workspace.register_artifact("project_relations", path, label="Project and subproject links")


def link_subproject(parent: SugarWorkspace, child: SugarWorkspace) -> dict[str, Any]:
    """Link an imported or existing workspace as a child without moving its files."""
    if parent.root == child.root or parent.manifest.project_id == child.manifest.project_id:
        raise ValueError("A project cannot be linked as its own subproject.")
    try:
        child_path = child.root.relative_to(parent.root).as_posix()
    except ValueError as exc:
        raise ValueError("A subproject must be stored inside its parent project's folder.") from exc
    parent_rel = _read_relation(parent)
    children = parent_rel.setdefault("children", [])
    if any(item.get("project_id") == child.manifest.project_id for item in children):
        return next(item for item in children if item.get("project_id") == child.manifest.project_id)
    child_rel = _read_relation(child)
    existing_parent = child_rel.get("parent")
    if existing_parent and existing_parent.get("project_id") not in {None, parent.manifest.project_id}:
        raise ValueError("The imported project already belongs to a different parent project.")
    relation = {"project_id": child.manifest.project_id, "name": child.manifest.name,
                "path": child_path, "created_at": utc_now()}
    children.append(relation)
    _write_relation(parent, parent_rel)
    child_rel.update({"schema_version": "1.0", "parent": {"project_id": parent.manifest.project_id,
                     "path_from_parent": child_path}})
    _write_relation(child, child_rel)
    log_project_event(parent, "subproject_linked", relation)
    return relation


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-") or "research-track"


def create_subproject(
    parent: SugarWorkspace,
    name: str,
    *,
    description: str = "",
    project_id: str | None = None,
) -> SugarWorkspace:
    clean_name = str(name or "").strip()
    if not clean_name:
        raise ValueError("A subproject name is required.")
    root = parent.root / "subprojects"
    root.mkdir(parents=True, exist_ok=True)
    target = root / _slug(clean_name)
    if target.exists():
        target = root / f"{_slug(clean_name)}-{uuid.uuid4().hex[:6]}"
    child = SugarWorkspace.create(target, name=clean_name, description=description, project_id=project_id)
    child_rel = {
        "schema_version": "1.0", "parent": {"project_id": parent.manifest.project_id,
                                                  "path_from_parent": child.root.relative_to(parent.root).as_posix()},
        "children": [], "created_at": utc_now(),
    }
    _write_relation(child, child_rel)
    parent_rel = _read_relation(parent)
    parent_rel.setdefault("children", []).append({"project_id": child.manifest.project_id,
                                                   "name": child.manifest.name,
                                                   "path": child.root.relative_to(parent.root).as_posix(),
                                                   "created_at": utc_now()})
    _write_relation(parent, parent_rel)
    log_project_event(parent, "subproject_created", {"project_id": child.manifest.project_id,
                                                       "name": clean_name, "path": child_rel["parent"]["path_from_parent"]})
    return child


def _project_row(path: Path) -> dict[str, Any] | None:
    try:
        workspace = SugarWorkspace.open(path)
    except (OSError, ValueError, json.JSONDecodeError):
        return None
    relation = _read_relation(workspace)
    status = workspace.status()
    history = list_project_history(workspace)
    successful = next((row.get("details", {}) for row in history
                       if row.get("event_type") == "research_run" and row.get("details", {}).get("status") == "succeeded"), {})
    issue_count = sum(1 for row in history if row.get("event_type") == "research_run"
                      and row.get("details", {}).get("status") in {"failed", "partial"})
    pending_review = 0
    for artifact in workspace.list_artifacts("state_assessments"):
        artifact_path = workspace.artifact_absolute_path(artifact)
        if not artifact_path.is_file() or artifact_path.suffix.casefold() not in {".jsonl", ".ndjson"}:
            continue
        try:
            with artifact_path.open("r", encoding="utf-8-sig") as stream:
                for line in stream:
                    if line.strip():
                        item = json.loads(line)
                        pending_review += int(item.get("review_state", "unreviewed") not in {"human_verified", "rejected"})
        except (OSError, ValueError):
            issue_count += 1
    return {
        "project_id": workspace.manifest.project_id, "name": workspace.manifest.name,
        "description": workspace.manifest.description, "root": str(workspace.root),
        "created_at": workspace.manifest.created_at, "updated_at": workspace.manifest.updated_at,
        "parent_project_id": (relation.get("parent") or {}).get("project_id", ""),
        "child_count": len(relation.get("children") or []), "artifact_count": status["artifact_count"],
        "missing_artifacts": status["missing_artifacts"], "pending_review": pending_review,
        "collection_issues": issue_count, "latest_successful_run": successful,
        "last_activity": history[0].get("occurred_at") if history else workspace.manifest.updated_at,
    }


def list_projects(root: str | Path, *, limit: int = 500) -> list[dict[str, Any]]:
    start = Path(root).expanduser().resolve()
    if start.is_file():
        start = start.parent
    if not start.exists():
        raise FileNotFoundError(start)
    candidates: list[Path] = []
    for current, directories, files in os.walk(start, followlinks=False):
        directories[:] = [name for name in directories if name not in {".git", ".sugar", "__pycache__", "node_modules"} and not name.startswith(".")]
        if MANIFEST_FILENAME in files:
            candidates.append(Path(current))
            directories[:] = [name for name in directories if name == "subprojects"]
        if len(candidates) >= max(1, min(5000, limit)):
            break
    rows = [row for path in candidates if (row := _project_row(path)) is not None]
    return sorted(rows, key=lambda row: (str(row.get("last_activity") or ""), row["name"].casefold()), reverse=True)


def dashboard(workspace: SugarWorkspace) -> dict[str, Any]:
    row = _project_row(workspace.root)
    if row is None:
        raise ValueError("Could not inspect the selected project.")
    row["artifacts"] = [
        {"kind": item.kind, "path": item.path, "label": item.label, "exists": item.exists,
         "updated_at": item.updated_at, "metadata": _safe_value(item.metadata)}
        for item in workspace.list_artifacts()
    ]
    row["children"] = _read_relation(workspace).get("children", [])
    row["recent_history"] = list_project_history(workspace, limit=30)
    return row


def log_search_plan_change(
    workspace: SugarWorkspace,
    *,
    branch_id: str,
    query: str = "",
    status: str = "",
    rationale: str = "",
    actor: str = "analyst",
    reason: str = "",
) -> dict[str, Any]:
    return log_project_event(workspace, "search_plan_change", {
        "branch_id": branch_id, "query": query, "status": status,
        "rationale": rationale, "reason": reason,
    }, actor=actor)
