from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import PostRecord
from .utils import safe_cell, utc_iso

PREFERRED_COLUMNS = [
    "platform", "native_id", "record_key", "content_type", "parent_record_key", "thread_root_key",
    "conversation_id", "published_at", "author_handle", "author_name", "author_location",
    "platform_language", "detected_language", "inferred_location", "location_confidence",
    "location_source", "location_reason", "latitude", "longitude", "original_text", "translated_text",
    "canonical_url", "engagement", "query_matches", "query", "source_mode", "source_host", "source_url",
    "collected_at", "collector_version", "schema_version", "raw_stats", "is_repost", "tweet_id",
    "post_url", "x_url", "username", "display_name", "date_iso", "date_raw", "translated_en", "is_retweet",
]


def records_to_frame(records: Iterable[PostRecord]) -> pd.DataFrame:
    df = pd.DataFrame([r.export_dict() for r in records])
    if df.empty:
        return df
    numeric = {"latitude", "longitude", "location_confidence"}
    for column in df.columns:
        if column not in numeric:
            df[column] = df[column].map(lambda x: safe_cell(x, formula_safe=True))
    ordered = [x for x in PREFERRED_COLUMNS if x in df.columns]
    ordered += [x for x in df.columns if x not in ordered]
    return df[ordered]


def save_records(records: Iterable[PostRecord], output_file: str | Path, *, metadata: dict | None = None) -> pd.DataFrame:
    records = list(records)
    if not records:
        raise ValueError("No records were collected.")
    df = records_to_frame(records)
    output_file = Path(output_file)
    csv_path = output_file if output_file.suffix.lower() == ".csv" else output_file.with_suffix(".csv")
    xlsx_path = csv_path.with_suffix(".xlsx")
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(csv_path, index=False, encoding="utf-8-sig", quoting=csv.QUOTE_ALL, lineterminator="\n")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="posts")
        ws = writer.sheets["posts"]
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        fill = PatternFill("solid", fgColor="D9EAF7"); font = Font(bold=True)
        for cell in ws[1]:
            cell.fill = fill; cell.font = font; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for idx, name in enumerate(df.columns, 1):
            width = 18
            if name in {"original_text", "translated_text", "translated_en", "location_reason"}: width = 60
            elif name in {"canonical_url", "source_url", "post_url", "x_url"}: width = 45
            ws.column_dimensions[get_column_letter(idx)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    metadata_path = csv_path.with_suffix(".metadata.json")
    payload = {
        "generated_at": utc_iso(), "records": len(df), "csv": csv_path.name, "xlsx": xlsx_path.name,
        **(metadata or {}),
    }
    metadata_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return df


def load_results(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    if not path.is_file(): raise FileNotFoundError(path)
    if path.suffix.lower() == ".xlsx": return pd.read_excel(path, sheet_name="posts")
    if path.suffix.lower() == ".csv": return pd.read_csv(path)
    raise ValueError("Results file must be CSV or XLSX.")
