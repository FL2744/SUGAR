from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .observations import OBSERVATION_SCHEMA_VERSION, ResearchObservation
from .utils import atomic_path, atomic_write_text, runtime_metadata, safe_cell, utc_iso

PREFERRED_OBSERVATION_COLUMNS = [
    "observation_id",
    "observation_type",
    "title",
    "summary",
    "observed_at",
    "activity_status",
    "location_label",
    "country",
    "region",
    "city",
    "latitude",
    "longitude",
    "location_basis",
    "location_confidence",
    "locations",
    "institution_name",
    "program_name",
    "actors",
    "audiences",
    "themes",
    "us_overlap",
    "overlap_note",
    "spatial_matches",
    "relevance",
    "relevance_confidence",
    "triage_labels",
    "triage_evidence",
    "ai_confidence",
    "ai_model",
    "ai_reason",
    "verification_state",
    "reviewer",
    "reviewed_at",
    "verification_notes",
    "primary_source_url",
    "evidence",
    "source_record_keys",
    "provenance",
    "created_at",
    "updated_at",
    "schema_version",
]

_NUMERIC_COLUMNS = {
    "latitude",
    "longitude",
    "location_confidence",
    "relevance_confidence",
    "ai_confidence",
}
_LONG_TEXT_COLUMNS = {
    "summary",
    "overlap_note",
    "locations",
    "spatial_matches",
    "triage_evidence",
    "ai_reason",
    "verification_notes",
    "evidence",
    "provenance",
}
_IDENTIFIER_COLUMNS = {"observation_id", "source_record_keys", "primary_source_url"}


def observations_to_frame(observations: Iterable[ResearchObservation]) -> pd.DataFrame:
    frame = pd.DataFrame([observation.export_dict() for observation in observations])
    if frame.empty:
        return frame
    frame["schema_version"] = OBSERVATION_SCHEMA_VERSION
    for column in frame.columns:
        if column not in _NUMERIC_COLUMNS:
            frame[column] = frame[column].map(lambda value: safe_cell(value, formula_safe=True))
    ordered = [column for column in PREFERRED_OBSERVATION_COLUMNS if column in frame.columns]
    ordered.extend(column for column in frame.columns if column not in ordered)
    return frame[ordered]


def save_observations(
    observations: Iterable[ResearchObservation],
    output_file: str | Path,
    *,
    metadata: dict | None = None,
) -> pd.DataFrame:
    observations = list(observations)
    if not observations:
        raise ValueError("No research observations were provided.")

    frame = observations_to_frame(observations)
    output_file = Path(output_file)
    csv_path = output_file if output_file.suffix.lower() == ".csv" else output_file.with_suffix(".csv")
    xlsx_path = csv_path.with_suffix(".xlsx")
    with atomic_path(csv_path) as temporary_csv:
        frame.to_csv(
            temporary_csv,
            index=False,
            encoding="utf-8-sig",
            quoting=csv.QUOTE_ALL,
            lineterminator="\n",
        )

    with atomic_path(xlsx_path) as temporary_xlsx:
        with pd.ExcelWriter(temporary_xlsx, engine="openpyxl") as writer:
            frame.to_excel(writer, index=False, sheet_name="observations")
            worksheet = writer.sheets["observations"]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            fill = PatternFill("solid", fgColor="D9EAF7")
            font = Font(bold=True)
            for cell in worksheet[1]:
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for index, name in enumerate(frame.columns, 1):
                width = 20
                if name in _LONG_TEXT_COLUMNS:
                    width = 60
                elif name in {"primary_source_url"}:
                    width = 45
                elif name in {"actors", "audiences", "themes", "us_overlap", "triage_labels"}:
                    width = 35
                worksheet.column_dimensions[get_column_letter(index)].width = width
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    metadata_path = csv_path.with_suffix(".metadata.json")
    payload = {
        "generated_at": utc_iso(),
        "sugar_version": __version__,
        "runtime": runtime_metadata(),
        "records": len(frame),
        "dataset_type": "research_observations",
        "observation_schema_version": OBSERVATION_SCHEMA_VERSION,
        "csv": csv_path.name,
        "xlsx": xlsx_path.name,
        **(metadata or {}),
    }
    atomic_write_text(metadata_path, json.dumps(payload, ensure_ascii=False, indent=2))
    return frame


def _load_jsonl_frame(path: Path) -> pd.DataFrame:
    rows: list[dict] = []
    with path.open("r", encoding="utf-8-sig") as stream:
        for line_number, line in enumerate(stream, 1):
            text = line.strip()
            if not text:
                continue
            try:
                payload = json.loads(text)
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSONL observation record on line {line_number}: {exc.msg}") from exc
            if not isinstance(payload, dict):
                raise ValueError(f"JSONL observation record on line {line_number} must be an object.")
            rows.append(payload)
    if not rows:
        raise ValueError("Observation JSONL file contains no records.")
    return pd.DataFrame(rows)


def load_observation_frame(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, dtype={column: str for column in _IDENTIFIER_COLUMNS})
    if suffix == ".xlsx":
        return pd.read_excel(path, sheet_name="observations", dtype={column: str for column in _IDENTIFIER_COLUMNS})
    if suffix in {".jsonl", ".ndjson"}:
        return _load_jsonl_frame(path)
    raise ValueError("Observation dataset must be CSV, XLSX, JSONL, or NDJSON.")


def load_observations(path: str | Path) -> list[ResearchObservation]:
    frame = load_observation_frame(path)
    observations: list[ResearchObservation] = []
    for raw in frame.to_dict(orient="records"):
        cleaned = {
            key: ("" if pd.isna(value) else value)
            for key, value in raw.items()
            if key in ResearchObservation.__dataclass_fields__ or key == "primary_source_url"
        }
        observations.append(ResearchObservation.from_export_dict(cleaned))
    return observations
