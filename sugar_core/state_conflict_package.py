from __future__ import annotations

import json
from dataclasses import asdict
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook

from .observation_storage import load_observations
from .observations import ResearchObservation
from .source_conflicts import (
    SourceConflict,
    load_source_conflicts,
    save_source_conflicts,
    source_conflict_summary,
)
from .state_review import export_review_workbook
from .state_schema import StateAssessment, USPresenceSite
from .state_workflow import (
    blank_state_assessments,
    build_review_queue,
    load_state_assessments,
    load_us_presence_sites,
    save_state_package,
)
from .utils import safe_cell

_CONFLICT_SECTION_START = "<!-- SUGAR_SOURCE_CONFLICTS_START -->"
_CONFLICT_SECTION_END = "<!-- SUGAR_SOURCE_CONFLICTS_END -->"


def _clean(value: Any) -> str:
    return " ".join(str(value or "").split())


def _stem(name: str) -> str:
    return "_".join(_clean(name).split()) or "state_research"


def _normalize_conflicts(
    conflicts: Iterable[SourceConflict | dict[str, Any]],
) -> list[SourceConflict]:
    return [row if isinstance(row, SourceConflict) else SourceConflict(**dict(row)) for row in conflicts]


def _formula_safe_frame(rows: Iterable[dict[str, Any]]) -> pd.DataFrame:
    frame = pd.DataFrame(list(rows))
    if not frame.empty:
        for column in frame.columns:
            frame[column] = frame[column].map(lambda value: safe_cell(value, formula_safe=True))
    return frame


def _observation_source_urls(observation: ResearchObservation | None) -> set[str]:
    if observation is None:
        return set()
    urls = {_clean(item.url) for item in observation.evidence if _clean(item.url)}
    urls.update(
        _clean(value) for value in observation.source_record_keys if _clean(value).startswith(("http://", "https://"))
    )
    return urls


def _conflict_applies_to_record(
    conflict: SourceConflict,
    observation: ResearchObservation | None,
    assessment: StateAssessment,
) -> bool:
    """Link a conflict only through source provenance that materially contributed to the record.

    Direct observation evidence always counts. Structured U.S.-service sources also count, except
    that service-topology conflicts require a program/service contribution. Audience-only service
    association is intentionally insufficient so audience similarity cannot manufacture a direct
    service-topology dependency.
    """
    conflict_urls = {claim.source_url for claim in conflict.claims}
    if conflict_urls & _observation_source_urls(observation):
        return True

    for source in assessment.us_overlap.service_sources:
        if not source.source_url or source.source_url not in conflict_urls:
            continue
        if conflict.conflict_type == "service_topology" and not source.program_service_matches:
            continue
        return True
    return False


def source_conflicts_for_record(
    observation: ResearchObservation | None,
    assessment: StateAssessment,
    source_conflicts: Iterable[SourceConflict | dict[str, Any]],
) -> list[SourceConflict]:
    """Return conflicts linked by exact, material source provenance, not topic similarity."""
    return [
        conflict
        for conflict in _normalize_conflicts(source_conflicts)
        if _conflict_applies_to_record(conflict, observation, assessment)
    ]


def build_conflict_aware_review_queue(
    observations: Iterable[ResearchObservation],
    assessments: Iterable[StateAssessment],
    source_conflicts: Iterable[SourceConflict | dict[str, Any]],
) -> list[dict[str, Any]]:
    observations = list(observations)
    assessments = list(assessments)
    conflicts = _normalize_conflicts(source_conflicts)
    observation_map = {row.observation_id: row for row in observations}
    assessment_map = {row.assessment_id: row for row in assessments}
    rows = build_review_queue(observations, assessments)

    for row in rows:
        assessment = assessment_map.get(str(row.get("assessment_id") or ""))
        if assessment is None:
            continue
        observation = observation_map.get(assessment.observation_id)
        matched = [conflict for conflict in conflicts if _conflict_applies_to_record(conflict, observation, assessment)]
        unresolved = [conflict for conflict in matched if conflict.requires_human_review]
        row["source_conflict_count"] = len(matched)
        row["source_conflicts_requiring_human_review"] = len(unresolved)
        row["source_conflict_ids"] = "; ".join(conflict.conflict_id for conflict in matched)
        row["source_conflict_topics"] = "; ".join(conflict.topic for conflict in matched)
        row["source_conflict_statuses"] = "; ".join(conflict.status for conflict in matched)
        if unresolved:
            row["review_priority"] = int(row.get("review_priority") or 0) + 4
            reasons = [value.strip() for value in str(row.get("reasons") or "").split(";") if value.strip()]
            reason = "source conflict affecting record requires human review"
            if reason not in reasons:
                reasons.append(reason)
            row["reasons"] = "; ".join(reasons)

    return sorted(
        rows,
        key=lambda row: (
            -int(row.get("review_priority") or 0),
            str(row.get("country") or ""),
            str(row.get("title") or ""),
        ),
    )


def _source_conflict_rows(conflicts: Iterable[SourceConflict]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for conflict in conflicts:
        preferred = conflict.preferred_claim
        rows.append(
            {
                "conflict_id": conflict.conflict_id,
                "topic": conflict.topic,
                "conflict_type": conflict.conflict_type,
                "status": conflict.status,
                "requires_human_review": conflict.requires_human_review,
                "is_resolved": conflict.is_resolved,
                "preferred_claim_id": conflict.preferred_claim_id,
                "preferred_source_url": preferred.source_url if preferred else "",
                "preferred_publisher": preferred.publisher if preferred else "",
                "preferred_authority_type": preferred.authority_type if preferred else "",
                "preferred_freshness": preferred.freshness if preferred else "",
                "preferred_effective_date": preferred.effective_date if preferred else "",
                "treatment": conflict.treatment,
                "preference_rationale": conflict.preference_rationale,
                "reviewer": conflict.reviewer,
                "review_note": conflict.review_note,
                "source_urls": "; ".join(claim.source_url for claim in conflict.claims),
                "claims_json": json.dumps(
                    [asdict(claim) for claim in conflict.claims],
                    ensure_ascii=False,
                    sort_keys=True,
                ),
            }
        )
    return rows


def render_source_conflict_section(
    source_conflicts: Iterable[SourceConflict | dict[str, Any]],
) -> str:
    conflicts = _normalize_conflicts(source_conflicts)
    summary = source_conflict_summary(conflicts)
    lines = [
        _CONFLICT_SECTION_START,
        "## Source Conflicts",
        "",
        (
            f"This package carries {summary['conflicts']} documented source conflict(s); "
            f"{summary['requiring_human_review']} still require human review. "
            "A provisional treatment is an operational representation of contradictory evidence, "
            "not a human adjudication and not evidence of influence, competition, or causal effect."
        ),
        "",
    ]
    for conflict in conflicts:
        preferred = conflict.preferred_claim
        preferred_label = preferred.publisher or preferred.source_label or preferred.source_url if preferred else "none"
        review = "human review required" if conflict.requires_human_review else "resolved"
        treatment = f" Treatment: {conflict.treatment}" if conflict.treatment else ""
        lines.append(
            f"- **{conflict.topic}** — {conflict.status}; {review}; preferred source: {preferred_label}.{treatment}"
        )
    lines.extend(["", _CONFLICT_SECTION_END, ""])
    return "\n".join(lines)


def _replace_conflict_section(text: str, section: str) -> str:
    start = text.find(_CONFLICT_SECTION_START)
    end = text.find(_CONFLICT_SECTION_END)
    if start >= 0 and end >= start:
        end += len(_CONFLICT_SECTION_END)
        text = text[:start].rstrip() + "\n\n" + text[end:].lstrip()
    return text.rstrip() + "\n\n" + section.rstrip() + "\n"


def _augment_audit(path: Path, conflicts: list[SourceConflict]) -> None:
    audit = json.loads(path.read_text(encoding="utf-8"))
    summary = source_conflict_summary(conflicts)
    audit["source_conflicts"] = summary
    if summary["requiring_human_review"]:
        finding = {
            "severity": "warning",
            "code": "source_conflict_requires_human_review",
            "assessment_id": "",
            "observation_id": "",
            "message": (
                f"{summary['requiring_human_review']} source conflict(s) remain open or provisional. "
                "Provisional treatments must not be represented as human-adjudicated fact."
            ),
        }
        findings = list(audit.get("findings") or [])
        if not any(item.get("code") == finding["code"] for item in findings):
            findings.append(finding)
            audit["warnings"] = int(audit.get("warnings") or 0) + 1
        audit["findings"] = findings
        if audit.get("status") == "pass":
            audit["status"] = "conditional"
    path.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def _conflict_fields_for_record(
    observation: ResearchObservation | None,
    assessment: StateAssessment,
    conflicts: list[SourceConflict],
) -> dict[str, Any]:
    matched = source_conflicts_for_record(observation, assessment, conflicts)
    unresolved = [conflict for conflict in matched if conflict.requires_human_review]
    return {
        "source_conflict_count": len(matched),
        "source_conflicts_requiring_human_review": len(unresolved),
        "source_conflict_ids": "; ".join(conflict.conflict_id for conflict in matched),
        "source_conflict_topics": "; ".join(conflict.topic for conflict in matched),
        "source_conflict_statuses": "; ".join(conflict.status for conflict in matched),
    }


def export_review_workbook_with_conflicts(
    observations: Iterable[ResearchObservation],
    assessments: Iterable[StateAssessment],
    output_file: str | Path,
    *,
    source_conflicts: Iterable[SourceConflict | dict[str, Any]] = (),
) -> str:
    """Export the analyst review workbook and carry source conflicts into the same file."""
    observations = list(observations)
    assessments = list(assessments)
    conflicts = _normalize_conflicts(source_conflicts)
    output = export_review_workbook(observations, assessments, output_file)
    if not conflicts:
        return output

    observation_map = {row.observation_id: row for row in observations}
    assessment_map = {row.assessment_id: row for row in assessments}
    workbook = load_workbook(output)
    assessment_sheet = workbook["assessments"]
    headers = {str(cell.value): cell.column for cell in assessment_sheet[1]}
    conflict_headers = [
        "source_conflict_count",
        "source_conflicts_requiring_human_review",
        "source_conflict_ids",
        "source_conflict_topics",
        "source_conflict_statuses",
    ]
    start_column = assessment_sheet.max_column + 1
    for offset, header in enumerate(conflict_headers):
        assessment_sheet.cell(row=1, column=start_column + offset, value=header)

    assessment_id_column = headers["assessment_id"]
    for row_number in range(2, assessment_sheet.max_row + 1):
        assessment_id = _clean(assessment_sheet.cell(row=row_number, column=assessment_id_column).value)
        assessment = assessment_map.get(assessment_id)
        if assessment is None:
            continue
        fields = _conflict_fields_for_record(
            observation_map.get(assessment.observation_id),
            assessment,
            conflicts,
        )
        for offset, header in enumerate(conflict_headers):
            assessment_sheet.cell(
                row=row_number,
                column=start_column + offset,
                value=safe_cell(fields[header], formula_safe=True),
            )

    if "source_conflicts" in workbook.sheetnames:
        del workbook["source_conflicts"]
    conflict_sheet = workbook.create_sheet("source_conflicts")
    conflict_frame = _formula_safe_frame(_source_conflict_rows(conflicts))
    if not conflict_frame.empty:
        conflict_sheet.append(list(conflict_frame.columns))
        for row in conflict_frame.itertuples(index=False, name=None):
            conflict_sheet.append(list(row))
        conflict_sheet.freeze_panes = "A2"
        conflict_sheet.auto_filter.ref = conflict_sheet.dimensions

    instructions = workbook["instructions"]
    instructions.append(
        [
            "Source conflicts",
            "Open or provisional source conflicts require human review. A provisional preferred claim is not a human adjudication and does not establish influence, competition, displacement, persuasion, or causal effect.",
        ]
    )
    workbook.save(output)
    return str(Path(output).resolve())


def augment_state_package_with_conflicts(
    observations: Iterable[ResearchObservation],
    assessments: Iterable[StateAssessment],
    output_directory: str | Path,
    *,
    name: str = "state_research",
    source_conflicts: Iterable[SourceConflict | dict[str, Any]],
) -> list[str]:
    """Attach source-conflict provenance to an already-created State package."""
    observations = list(observations)
    assessments = list(assessments)
    conflicts = _normalize_conflicts(source_conflicts)
    if not conflicts:
        return []

    out_dir = Path(output_directory).expanduser().resolve()
    stem = _stem(name)
    conflict_path = out_dir / f"{stem}.source_conflicts.json"
    queue_path = out_dir / f"{stem}.review_queue.csv"
    xlsx_path = out_dir / f"{stem}.state.xlsx"
    audit_path = out_dir / f"{stem}.audit.json"
    brief_path = out_dir / f"{stem}.brief.md"
    snapshot_path = out_dir / f"{stem}.snapshot.json"

    save_source_conflicts(conflicts, conflict_path)

    review_rows = build_conflict_aware_review_queue(observations, assessments, conflicts)
    review_frame = _formula_safe_frame(review_rows)
    conflict_frame = _formula_safe_frame(_source_conflict_rows(conflicts))
    review_frame.to_csv(queue_path, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        review_frame.to_excel(writer, index=False, sheet_name="review_queue")
        conflict_frame.to_excel(writer, index=False, sheet_name="source_conflicts")

    _augment_audit(audit_path, conflicts)

    brief = brief_path.read_text(encoding="utf-8")
    brief_path.write_text(
        _replace_conflict_section(brief, render_source_conflict_section(conflicts)),
        encoding="utf-8",
    )

    package_audit = json.loads(audit_path.read_text(encoding="utf-8"))
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    snapshot["audit_status"] = package_audit.get("status")
    snapshot["source_conflicts"] = source_conflict_summary(conflicts)
    snapshot["source_conflict_ids"] = [conflict.conflict_id for conflict in conflicts]
    snapshot_path.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    return [str(conflict_path.resolve())]


def save_state_package_with_conflicts(
    observations: Iterable[ResearchObservation],
    assessments: Iterable[StateAssessment],
    output_directory: str | Path,
    *,
    name: str = "state_research",
    us_sites: Iterable[USPresenceSite] = (),
    previous_assessments: Iterable[StateAssessment] | None = None,
    title: str = "PRC Cultural Influence Network Research Update",
    source_conflicts: Iterable[SourceConflict | dict[str, Any]] = (),
) -> list[str]:
    observations = list(observations)
    assessments = list(assessments)
    sites = list(us_sites)
    conflicts = _normalize_conflicts(source_conflicts)
    outputs = save_state_package(
        observations,
        assessments,
        output_directory,
        name=name,
        us_sites=sites,
        previous_assessments=previous_assessments,
        title=title,
    )
    outputs.extend(
        augment_state_package_with_conflicts(
            observations,
            assessments,
            output_directory,
            name=name,
            source_conflicts=conflicts,
        )
    )
    return outputs


def package_from_files_with_conflicts(
    observations_file: str | Path,
    output_directory: str | Path,
    *,
    assessments_file: str | Path | None = None,
    us_sites_file: str | Path | None = None,
    previous_assessments_file: str | Path | None = None,
    source_conflicts_file: str | Path | None = None,
    name: str = "state_research",
    title: str = "PRC Cultural Influence Network Research Update",
) -> list[str]:
    observations = load_observations(observations_file)
    assessments = (
        load_state_assessments(assessments_file) if assessments_file else blank_state_assessments(observations)
    )
    sites = load_us_presence_sites(us_sites_file) if us_sites_file else []
    previous = load_state_assessments(previous_assessments_file) if previous_assessments_file else None
    conflicts = load_source_conflicts(source_conflicts_file) if source_conflicts_file else []
    return save_state_package_with_conflicts(
        observations,
        assessments,
        output_directory,
        name=name,
        us_sites=sites,
        previous_assessments=previous,
        title=title,
        source_conflicts=conflicts,
    )
