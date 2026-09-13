from __future__ import annotations

import json
import math
import re
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd

from .storage import load_results

ENGAGEMENT_KEYS = ("likes", "replies", "reposts", "quotes", "bookmarks")


def _json_dict(value: Any) -> dict:
    if isinstance(value, dict): return value
    try:
        parsed = json.loads(str(value))
        return parsed if isinstance(parsed, dict) else {}
    except Exception:
        return {}


def _prepare(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, Any]]:
    if df.empty: raise ValueError("The selected results file contains no records.")
    work = df.copy()
    engagement_source = work.get("engagement")
    if engagement_source is None:
        engagement_source = work.get("raw_stats", pd.Series(["{}"] * len(work), index=work.index))
    stats = engagement_source.map(_json_dict)

    def canonical(d: dict, key: str) -> int:
        aliases = {
            "likes": ("likes", "like_count", "favourite_count"),
            "replies": ("replies", "reply_count"),
            "reposts": ("reposts", "retweet_count", "repost_count", "reblog_count"),
            "quotes": ("quotes", "quote_count"),
            "bookmarks": ("bookmarks", "bookmark_count"),
            "impressions": ("impressions", "impression_count"),
        }[key]
        for alias in aliases:
            if alias in d:
                try: return int(float(d.get(alias, 0) or 0))
                except Exception: return 0
        return 0

    for key in (*ENGAGEMENT_KEYS, "impressions"):
        work[key] = stats.map(lambda d, k=key: canonical(d, k))
    work["engagement_total"] = work[list(ENGAGEMENT_KEYS)].sum(axis=1)
    work["engagement_rate"] = work.apply(
        lambda r: r.engagement_total / r.impressions if r.impressions else math.nan, axis=1
    )
    date_col = "published_at" if "published_at" in work else "date_iso"
    work["date"] = pd.to_datetime(work.get(date_col), errors="coerce", utc=True)
    platform_col = work.get("platform", pd.Series(["unknown"] * len(work), index=work.index))
    work["platform"] = platform_col.fillna("unknown").replace("", "unknown")
    lang_col = "detected_language" if "detected_language" in work else "platform_language"
    work["language"] = work.get(lang_col, pd.Series(["unknown"] * len(work), index=work.index)).fillna("unknown").replace("", "unknown")
    loc = work.get("inferred_location", pd.Series(["Unassigned"] * len(work), index=work.index))
    work["location"] = loc.fillna("Unassigned").replace("", "Unassigned")
    id_col = "native_id" if "native_id" in work else "tweet_id"
    ids = work.get(id_col, pd.Series([""] * len(work), index=work.index)).fillna("").astype(str).str.strip()
    identity = pd.DataFrame({"platform": work["platform"].astype(str), "native_id": ids}, index=work.index)
    unique_ids = len(identity.loc[identity.native_id.ne("")].drop_duplicates())
    valid_dates = work.date.dropna()
    return work, {
        "records": len(work), "unique_ids": int(unique_ids),
        "engagement": int(work.engagement_total.sum()), "impressions": int(work.impressions.sum()),
        "start": valid_dates.min().strftime("%Y-%m-%d") if len(valid_dates) else "Unknown",
        "end": valid_dates.max().strftime("%Y-%m-%d") if len(valid_dates) else "Unknown",
    }


def _terms(work: pd.DataFrame) -> list[tuple[str, int]]:
    text_col = "translated_text" if "translated_text" in work else "translated_en"
    text = " ".join(work.get(text_col, pd.Series(dtype=str)).fillna("").astype(str))
    if not text.strip(): text = " ".join(work.get("original_text", pd.Series(dtype=str)).fillna("").astype(str))
    stop = {"the","and","that","with","from","this","have","for","are","was","were","their","about","http","https","www","com"}
    words = re.findall(r"[A-Za-z][A-Za-z'-]{3,}", text.casefold())
    return Counter(w for w in words if w not in stop).most_common(15)


def _bar(series: pd.Series, title: str, path: Path) -> None:
    values = series.head(10).sort_values()
    fig, ax = plt.subplots(figsize=(8, 3.5))
    ax.barh([str(x) for x in values.index], values.values)
    ax.set_title(title); ax.set_xlabel("Records"); ax.grid(axis="x", alpha=.2)
    fig.tight_layout(); fig.savefig(path, dpi=160, bbox_inches="tight"); plt.close(fig)


def _docx(work: pd.DataFrame, metrics: dict, charts: dict[str, Path], source: str, output: str) -> None:
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.7); sec.left_margin = sec.right_margin = Inches(.75)
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("SUGAR Collection Analysis", 0)
    doc.add_paragraph(f"Deterministic descriptive review of {metrics['records']:,} collected records.")
    table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"
    for i, h in enumerate(("Measure","Value")): table.rows[0].cells[i].text = h
    for key, value in [
        ("Records", f"{metrics['records']:,}"), ("Unique IDs", f"{metrics['unique_ids']:,}"),
        ("Observed dates", f"{metrics['start']} to {metrics['end']}"),
        ("Recorded engagement", f"{metrics['engagement']:,}"), ("Recorded impressions", f"{metrics['impressions']:,}"),
    ]:
        cells = table.add_row().cells; cells[0].text = key; cells[1].text = value
    doc.add_heading("Collection composition", 1)
    doc.add_picture(str(charts["platform"]), width=Inches(6.5)); doc.add_picture(str(charts["language"]), width=Inches(6.5))
    doc.add_heading("Geographic distribution", 1); doc.add_picture(str(charts["location"]), width=Inches(6.5))
    doc.add_paragraph("Locations are model-assisted or profile-derived research fields, not verified precise geotags. Use source evidence and human review before drawing geographic conclusions.")
    doc.add_heading("Engagement", 1)
    platform_eng = work.groupby("platform")["engagement_total"].sum().sort_values(ascending=False)
    for platform, value in platform_eng.items(): doc.add_paragraph(f"{platform}: {int(value):,} recorded interactions", style="List Bullet")
    doc.add_heading("Recurring vocabulary", 1)
    doc.add_paragraph(", ".join(f"{w} ({c})" for w,c in _terms(work)) or "No recurring terms calculated.")
    doc.add_heading("Method caveats", 1)
    doc.add_paragraph("Counts describe retrieved records, not the full population of online discussion. Platform search coverage, pagination, source availability, and query design affect what is observed. Canonical engagement maps platform-specific likes/favorites, replies, reposts/retweets/reblogs, quotes, and bookmarks into common fields without inventing missing metrics.")
    doc.add_paragraph(f"Source file: {Path(source).name}")
    doc.save(output)


def _pdf(work: pd.DataFrame, metrics: dict, charts: dict[str, Path], source: str, output: str) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    styles = getSampleStyleSheet(); story = [Paragraph("SUGAR Collection Analysis", styles["Title"]), Spacer(1, 8)]
    data = [["Measure","Value"], ["Records",f"{metrics['records']:,}"], ["Unique IDs",f"{metrics['unique_ids']:,}"],
            ["Observed dates",f"{metrics['start']} to {metrics['end']}"], ["Recorded engagement",f"{metrics['engagement']:,}"],
            ["Recorded impressions",f"{metrics['impressions']:,}"]]
    t = Table(data, colWidths=[2.2*inch, 3.8*inch]); t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.4,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.lightgrey)])); story += [t, Spacer(1,12)]
    for title, key in (("Collection composition","platform"),("Languages","language"),("Geographic distribution","location")):
        story += [Paragraph(title, styles["Heading1"]), Image(str(charts[key]), width=6.3*inch, height=2.8*inch), Spacer(1,8)]
    story += [PageBreak(), Paragraph("Method caveats", styles["Heading1"]), Paragraph(
        "Counts describe retrieved records, not the full population of online discussion. Platform search coverage, pagination, source availability, and query design affect what is observed. Locations require source review and human verification before geographic claims.", styles["BodyText"]),
        Spacer(1,8), Paragraph(f"Source file: {Path(source).name}", styles["BodyText"])]
    SimpleDocTemplate(output, pagesize=letter, leftMargin=.7*inch, rightMargin=.7*inch, topMargin=.7*inch, bottomMargin=.7*inch).build(story)


def create_analysis_report(source_file: str, output_stem: str, output_format: str = "both") -> list[str]:
    if output_format not in {"docx","pdf","both"}: raise ValueError("output_format must be docx, pdf, or both")
    work, metrics = _prepare(load_results(source_file)); base = Path(output_stem).expanduser().resolve(); base.parent.mkdir(parents=True, exist_ok=True)
    outputs: list[str] = []
    with tempfile.TemporaryDirectory(prefix="sugar-report-") as tmp:
        tmp = Path(tmp); charts = {k: tmp/f"{k}.png" for k in ("platform","language","location")}
        _bar(work.platform.value_counts(), "Records by platform", charts["platform"])
        _bar(work.language.value_counts(), "Records by detected language", charts["language"])
        _bar(work.location.value_counts(), "Top inferred locations", charts["location"])
        if output_format in {"docx","both"}: path=str(base.with_suffix(".docx")); _docx(work,metrics,charts,source_file,path); outputs.append(path)
        if output_format in {"pdf","both"}: path=str(base.with_suffix(".pdf")); _pdf(work,metrics,charts,source_file,path); outputs.append(path)
    return outputs